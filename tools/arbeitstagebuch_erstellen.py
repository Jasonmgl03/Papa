import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, DataBarRule, CellIsRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment

F = "Arial"
NAVY = "1F3A5F"; INPUT = "FFF8DC"; CALC = "F2F2F2"; LIGHT = "DCE6F1"
def font(**kw): return Font(name=F, **{"size": 10, **kw})
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_fill = PatternFill("solid", fgColor=NAVY)
def header(c):
    c.font = font(bold=True, color="FFFFFF"); c.fill = hdr_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = box

N = 2000            # Eingabezeilen
R0 = 5; R1 = R0 + N - 1
wb = Workbook()

# ---------------- Eingabe ----------------
ws = wb.active; ws.title = "Eingabe"
ws["A1"] = "Arbeitstagebuch – tägliche Erfassung"; ws["A1"].font = font(size=16, bold=True, color=NAVY)
ws["A2"] = ("Gelbe Spalten ausfüllen, graue rechnen automatisch.  Datum: Strg + .   |   Uhrzeit: Strg + Umschalt + .   |   "
            "Entweder 'Dauer' eintragen ODER Start + Ende.   Beispielzeilen einfach löschen.")
ws["A2"].font = font(italic=True, color="595959")
cols = [("Datum", 12, True), ("Projekt", 22, True), ("Kategorie", 16, True), ("Aufgabe / Tätigkeit", 40, True),
        ("Dauer\n(Std.)", 9, True), ("Start", 8, True), ("Ende", 8, True), ("Status", 11, True),
        ("Priorität", 10, True), ("Notiz", 28, True),
        ("Stunden", 9, False), ("Wochentag", 11, False), ("KW", 6, False), ("Monat", 9, False)]
for i, (name, w, inp) in enumerate(cols, 1):
    c = ws.cell(row=4, column=i, value=name); header(c)
    ws.column_dimensions[c.column_letter].width = w
ws.row_dimensions[4].height = 30
inp_fill = PatternFill("solid", fgColor=INPUT); calc_fill = PatternFill("solid", fgColor=CALC)
for r in range(R0, R1 + 1):
    ws[f"K{r}"] = f'=IF(E{r}<>"",E{r},IF(AND(F{r}<>"",G{r}<>""),MOD(G{r}-F{r},1)*24,""))'
    ws[f"L{r}"] = f'=IF(A{r}="","",CHOOSE(WEEKDAY(A{r},2),"Montag","Dienstag","Mittwoch","Donnerstag","Freitag","Samstag","Sonntag"))'
    ws[f"M{r}"] = f'=IF(A{r}="","",_xlfn.ISOWEEKNUM(A{r}))'
    ws[f"N{r}"] = f'=IF(A{r}="","",DATE(YEAR(A{r}),MONTH(A{r}),1))'
    for col in "ABCDEFGHIJ":
        c = ws[f"{col}{r}"]; c.fill = inp_fill; c.border = box; c.font = font()
    for col in "KLMN":
        c = ws[f"{col}{r}"]; c.fill = calc_fill; c.border = box; c.font = font()
    ws[f"A{r}"].number_format = "DD.MM.YYYY"
    ws[f"E{r}"].number_format = "0.00"; ws[f"K{r}"].number_format = "0.00"
    ws[f"F{r}"].number_format = "hh:mm"; ws[f"G{r}"].number_format = "hh:mm"
    ws[f"N{r}"].number_format = "MMM YYYY"
    ws[f"D{r}"].alignment = Alignment(wrap_text=False)
ws.freeze_panes = "B5"
ws.auto_filter.ref = f"A4:N{R1}"
ws.sheet_view.zoomScale = 110
ws["E4"].comment = Comment("Dauer in Stunden, z. B. 1,5 = 1 Std. 30 Min.\nLeer lassen, wenn Start und Ende eingetragen werden.", "Vorlage")

# Beispielzeilen (zum Löschen)
ex = [
    (dt.date(2026, 9, 28), "Projekt Alpha", "Planung", "Projektplan Q4 erstellen", 2.0, None, None, "Erledigt", "Hoch", ""),
    (dt.date(2026, 9, 28), "Verwaltung", "E-Mail / Orga", "Posteingang abarbeiten", 1.0, None, None, "Erledigt", "Niedrig", ""),
    (dt.date(2026, 9, 29), "Kunde Müller", "Besprechung", "Abstimmung Angebot", None, dt.time(9, 0), dt.time(10, 30), "Erledigt", "Mittel", "Rückruf Fr."),
    (dt.date(2026, 9, 30), "Projekt Alpha", "Umsetzung", "Konzept Teil 1", 4.5, None, None, "Erledigt", "Hoch", ""),
    (dt.date(2026, 10, 1), "Kunde Müller", "Umsetzung", "Angebot ausarbeiten", 3.0, None, None, "In Arbeit", "Hoch", ""),
    (dt.date(2026, 10, 2), "Weiterbildung", "Lernen", "Excel-Kurs Modul 2", 1.5, None, None, "Erledigt", "Niedrig", ""),
    (dt.date(2026, 10, 5), "Projekt Alpha", "Umsetzung", "Konzept Teil 2", None, dt.time(8, 0), dt.time(11, 15), "In Arbeit", "Hoch", ""),
    (dt.date(2026, 10, 5), "Verwaltung", "E-Mail / Orga", "Wochenplanung", 0.5, None, None, "Offen", "Mittel", "Beispielzeile – löschen"),
]
for i, row in enumerate(ex):
    for j, v in enumerate(row):
        if v is not None and v != "":
            ws.cell(row=R0 + i, column=j + 1, value=v)

# ---------------- Listen ----------------
ls = wb.create_sheet("Projekte")
ls["A1"] = "Projekte & Auswahllisten"; ls["A1"].font = font(size=16, bold=True, color=NAVY)
ls["A2"] = "Hier eigene Projekte eintragen (gelbe Felder). Sie erscheinen sofort in der Auswahl auf 'Eingabe' und im Dashboard."
ls["A2"].font = font(italic=True, color="595959")
for i, (h, w) in enumerate([("Projekt", 24), ("Budget\n(Std.)", 10), ("Deadline", 12), ("Bemerkung", 30)], 1):
    c = ls.cell(row=4, column=i, value=h); header(c); ls.column_dimensions[c.column_letter].width = w
ls.row_dimensions[4].height = 30
projects = [("Projekt Alpha", 80, dt.date(2026, 12, 18), "Beispiel"), ("Kunde Müller", 20, dt.date(2026, 10, 30), "Beispiel"),
            ("Verwaltung", None, None, "laufend"), ("Weiterbildung", 24, dt.date(2026, 12, 31), "")]
P0, P1 = 5, 34   # 30 Projekte
for r in range(P0, P1 + 1):
    for col in "ABCD":
        c = ls[f"{col}{r}"]; c.fill = inp_fill; c.border = box; c.font = font()
    ls[f"C{r}"].number_format = "DD.MM.YYYY"; ls[f"B{r}"].number_format = "0"
for i, p in enumerate(projects):
    for j, v in enumerate(p):
        if v is not None: ls.cell(row=P0 + i, column=j + 1, value=v)

def simple_list(col, title, values, width=18):
    c = ls[f"{col}4"]; c.value = title; header(c); ls.column_dimensions[col].width = width
    for i in range(15):
        cc = ls[f"{col}{5+i}"]; cc.fill = inp_fill; cc.border = box; cc.font = font()
        if i < len(values): cc.value = values[i]
simple_list("F", "Kategorie", ["Planung", "Umsetzung", "Besprechung", "Telefon", "E-Mail / Orga", "Recherche", "Dokumentation", "Fahrt", "Lernen", "Pause"])
simple_list("H", "Status", ["Offen", "In Arbeit", "Erledigt", "Wartet"], 12)
simple_list("J", "Priorität", ["Hoch", "Mittel", "Niedrig"], 12)
ls.column_dimensions["E"].width = 3; ls.column_dimensions["G"].width = 3; ls.column_dimensions["I"].width = 3
ls.freeze_panes = "A5"

for nm, ref in [("ListeProjekte", f"Projekte!$A${P0}:$A${P1}"), ("ListeKategorie", "Projekte!$F$5:$F$19"),
                ("ListeStatus", "Projekte!$H$5:$H$19"), ("ListePrio", "Projekte!$J$5:$J$19")]:
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)

def dv_list(rng, name, col):
    dv = DataValidation(type="list", formula1=f"={name}", allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv); dv.add(f"{col}{R0}:{col}{R1}")
dv_list(None, "ListeProjekte", "B"); dv_list(None, "ListeKategorie", "C")
dv_list(None, "ListeStatus", "H"); dv_list(None, "ListePrio", "I")
dvd = DataValidation(type="date", operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2040,12,31)",
                     showErrorMessage=True, errorTitle="Datum", error="Bitte ein gültiges Datum eingeben (z. B. Strg + .)")
ws.add_data_validation(dvd); dvd.add(f"A{R0}:A{R1}")
dvh = DataValidation(type="decimal", operator="between", formula1="0", formula2="24", showErrorMessage=True,
                     errorTitle="Dauer", error="Dauer in Stunden zwischen 0 und 24, z. B. 1,5")
ws.add_data_validation(dvh); dvh.add(f"E{R0}:E{R1}")

# Bedingte Formatierung Eingabe: heute hervorheben, Status
ws.conditional_formatting.add(f"A{R0}:A{R1}", FormulaRule(formula=[f"$A{R0}=TODAY()"], fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(bold=True)))
ws.conditional_formatting.add(f"H{R0}:H{R1}", CellIsRule(operator="equal", formula=['"Erledigt"'], font=Font(color="008000")))
ws.conditional_formatting.add(f"H{R0}:H{R1}", CellIsRule(operator="equal", formula=['"Offen"'], font=Font(color="C00000", bold=True)))
ws.conditional_formatting.add(f"I{R0}:I{R1}", CellIsRule(operator="equal", formula=['"Hoch"'], font=Font(color="C00000", bold=True)))

# ---------------- Dashboard ----------------
db = wb.create_sheet("Dashboard", 0)
db.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGHIJKL", [2, 22, 12, 12, 12, 12, 12, 3, 22, 12, 12, 12]):
    db.column_dimensions[col].width = w
db["B1"] = "Dashboard – Zeitplanung"; db["B1"].font = font(size=18, bold=True, color=NAVY)
E = lambda col: f"Eingabe!${col}${R0}:${col}${R1}"
STD, DAT, PRJ, KAT, STA, PRI = E("K"), E("A"), E("B"), E("C"), E("H"), E("I")

# Einstellungen
db["B3"] = "Stichtag"; db["C3"] = "=TODAY()"
db["B4"] = "Wochenziel (Std.)"; db["C4"] = 40
db["B5"] = "Arbeitstage / Woche"; db["C5"] = 5
for r in (3, 4, 5):
    db[f"B{r}"].font = font(bold=True)
    c = db[f"C{r}"]; c.fill = inp_fill; c.border = box; c.font = font(color="0000FF"); c.alignment = Alignment(horizontal="center")
db["C3"].number_format = "DD.MM.YYYY"
db["D3"] = "← heute (überschreibbar, z. B. für Rückblick)"; db["D4"] = "← Ihre Annahme, anpassen"
db["D3"].font = db["D4"].font = font(italic=True, size=9, color="808080")
# Hilfswerte
db["E5"] = "Wochenbeginn:"; db["F5"] = "=C3-WEEKDAY(C3,2)+1"; db["F5"].number_format = "DD.MM.YYYY"
db["G5"] = "=\"KW \"&_xlfn.ISOWEEKNUM(C3)"
for a in ("E5", "F5", "G5"): db[a].font = font(size=9, color="808080")

# KPI-Kacheln
kpis = [("Heute", f"=SUMIFS({STD},{DAT},$C$3)", "0.0\" h\""),
        ("Diese Woche", f"=SUMIFS({STD},{DAT},\">=\"&$F$5,{DAT},\"<=\"&$F$5+6)", "0.0\" h\""),
        ("Wochenziel erreicht", "=IF($C$4=0,0,C9/$C$4)", "0%"),
        ("Dieser Monat", f"=SUMIFS({STD},{DAT},\">=\"&DATE(YEAR($C$3),MONTH($C$3),1),{DAT},\"<=\"&EOMONTH($C$3,0))", "0.0\" h\""),
        ("Ø pro Arbeitstag (Woche)", f"=IFERROR(C9/MAX(1,MIN($C$5,COUNTIFS({DAT},\">=\"&$F$5,{DAT},\"<=\"&$C$3,{STD},\">0\")/1)),0)", "0.0\" h\""),
        ("Offene Aufgaben", f"=COUNTIFS({STA},\"Offen\")+COUNTIFS({STA},\"In Arbeit\")+COUNTIFS({STA},\"Wartet\")", "0"),
        ]
# Layout: zwei Reihen à 3 Kacheln (Label-Zeile + Wert-Zeile)
pos = [("C", 8), ("E", 8), ("G", 8), ("C", 11), ("E", 11), ("G", 11)]
pos = [("B", 7), ("D", 7), ("F", 7), ("B", 10), ("D", 10), ("F", 10)]
tile = PatternFill("solid", fgColor=LIGHT)
cells_of = {}
for (label, f, fmt), (col, r) in zip(kpis, pos):
    col2 = chr(ord(col) + 1) if col != "B" else "C"
    db.merge_cells(f"{col}{r}:{col2}{r}"); db.merge_cells(f"{col}{r+1}:{col2}{r+1}")
    l = db[f"{col}{r}"]; l.value = label; l.font = font(size=9, bold=True, color="595959")
    v = db[f"{col}{r+1}"]; v.value = f; v.number_format = fmt; v.font = font(size=20, bold=True, color=NAVY)
    for rr in (r, r + 1):
        for cc in (col, col2):
            db[f"{cc}{rr}"].fill = tile
            db[f"{cc}{rr}"].alignment = Alignment(horizontal="center", vertical="center")
    db.row_dimensions[r + 1].height = 32
    cells_of[label] = f"{col}{r+1}"
# Wochenziel-Formel auf Wochen-Kachel beziehen
db[cells_of["Wochenziel erreicht"]] = f"=IF($C$4=0,0,{cells_of['Diese Woche']}/$C$4)"
db[cells_of["Ø pro Arbeitstag (Woche)"]] = (f"=IFERROR({cells_of['Diese Woche']}/SUMPRODUCT(--(COUNTIFS({DAT},$F$5+ROW($1:$7)-1,{STD},\">0\")>0)),0)")
db.conditional_formatting.add(cells_of["Wochenziel erreicht"], CellIsRule(operator="greaterThanOrEqual", formula=["1"], font=Font(color="008000", bold=True, size=20)))

# Projektübersicht
r = 14
db[f"B{r}"] = "Projekte"; db[f"B{r}"].font = font(size=13, bold=True, color=NAVY)
heads = ["Projekt", "Woche (h)", "Monat (h)", "Gesamt (h)", "Budget (h)", "Verbraucht"]
for i, h in enumerate(heads):
    header(db.cell(row=r + 1, column=2 + i, value=h))
db.cell(row=r + 1, column=8)
PR0 = r + 2
for i in range(30):
    rr = PR0 + i; p = f"Projekte!$A${P0+i}"
    db[f"B{rr}"] = f'=IF({p}="","",{p})'
    db[f"C{rr}"] = f'=IF($B{rr}="","",SUMIFS({STD},{PRJ},$B{rr},{DAT},">="&$F$5,{DAT},"<="&$F$5+6))'
    db[f"D{rr}"] = f'=IF($B{rr}="","",SUMIFS({STD},{PRJ},$B{rr},{DAT},">="&DATE(YEAR($C$3),MONTH($C$3),1),{DAT},"<="&EOMONTH($C$3,0)))'
    db[f"E{rr}"] = f'=IF($B{rr}="","",SUMIFS({STD},{PRJ},$B{rr}))'
    db[f"F{rr}"] = f'=IF(OR($B{rr}="",Projekte!$B${P0+i}=""),"",Projekte!$B${P0+i})'
    db[f"G{rr}"] = f'=IF(OR($B{rr}="",F{rr}=""),"",E{rr}/F{rr})'
    for col in "BCDEFG":
        c = db[f"{col}{rr}"]; c.font = font()
    for col in "CDEF": db[f"{col}{rr}"].number_format = "0.0;-0.0;\"–\""
    db[f"G{rr}"].number_format = "0%"
PR1 = PR0 + 29
db.conditional_formatting.add(f"G{PR0}:G{PR1}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="5B9BD5"))
db.conditional_formatting.add(f"G{PR0}:G{PR1}", FormulaRule(formula=[f'AND(ISNUMBER(G{PR0}),G{PR0}>1)'], font=Font(color="C00000", bold=True)))
db.conditional_formatting.add(f"G{PR0}:G{PR0+29}", FormulaRule(formula=[f'AND(ISNUMBER(G{PR0}),G{PR0}>=0.8,G{PR0}<=1)'], font=Font(color="C65911", bold=True)))
# Leere Zeilen ausblenden-Optik: Rahmen nur wenn Projekt vorhanden
# Summe
db[f"B{PR1+1}"] = "Summe"; db[f"B{PR1+1}"].font = font(bold=True)
for col in "CDE":
    db[f"{col}{PR1+1}"] = f"=SUM({col}{PR0}:{col}{PR1})"; db[f"{col}{PR1+1}"].font = font(bold=True)
    db[f"{col}{PR1+1}"].number_format = "0.0"
    db[f"{col}{PR1+1}"].border = Border(top=Side(style="medium", color=NAVY))
db[f"B{PR1+1}"].border = Border(top=Side(style="medium", color=NAVY))
db[f"B{PR1+2}"] = "Ohne Projekt (h)"; db[f"E{PR1+2}"] = f'=SUM({STD})-E{PR1+1}'
db[f"B{PR1+2}"].font = font(size=9, italic=True, color="808080"); db[f"E{PR1+2}"].font = font(size=9, italic=True, color="808080")
db[f"E{PR1+2}"].number_format = "0.0"

# Rechte Spalte: Deadlines
db["I14"] = "Deadlines"; db["I14"].font = font(size=13, bold=True, color=NAVY)
for i, h in enumerate(["Projekt", "Deadline", "Tage noch", "Rest (h)"]):
    header(db.cell(row=15, column=9 + i, value=h))
for i in range(30):
    rr = 16 + i; pr = P0 + i
    db[f"I{rr}"] = f'=IF(Projekte!$C${pr}="","",Projekte!$A${pr})'
    db[f"J{rr}"] = f'=IF(I{rr}="","",Projekte!$C${pr})'
    db[f"K{rr}"] = f'=IF(I{rr}="","",J{rr}-$C$3)'
    db[f"L{rr}"] = f'=IF(OR(I{rr}="",Projekte!$B${pr}=""),"",Projekte!$B${pr}-SUMIFS({STD},{PRJ},Projekte!$A${pr}))'
    for col in "IJKL": db[f"{col}{rr}"].font = font()
    db[f"J{rr}"].number_format = "DD.MM.YYYY"; db[f"K{rr}"].number_format = "0"; db[f"L{rr}"].number_format = "0.0"
db.conditional_formatting.add("K16:K45", FormulaRule(formula=['AND(ISNUMBER(K16),K16<0)'], fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(color="9C0006", bold=True)))
db.conditional_formatting.add("K16:K45", FormulaRule(formula=['AND(ISNUMBER(K16),K16>=0,K16<=14)'], fill=PatternFill("solid", fgColor="FFEB9C"), font=Font(color="9C5700", bold=True)))
db.conditional_formatting.add("L16:L45", FormulaRule(formula=['AND(ISNUMBER(L16),L16<0)'], font=Font(color="C00000", bold=True)))

# Offene Aufgaben nach Priorität
db["I48"] = "Offene Aufgaben nach Priorität"; db["I48"].font = font(size=13, bold=True, color=NAVY)
for i, h in enumerate(["Priorität", "Anzahl"]): header(db.cell(row=49, column=9 + i, value=h))
for i, p in enumerate(["Hoch", "Mittel", "Niedrig"]):
    rr = 50 + i; db[f"I{rr}"] = p; db[f"I{rr}"].font = font()
    db[f"J{rr}"] = f'=COUNTIFS({PRI},I{rr},{STA},"<>Erledigt",{DAT},"<>")'; db[f"J{rr}"].font = font()
    db[f"J{rr}"].alignment = Alignment(horizontal="center")
db["I50"].font = font(bold=True, color="C00000")

# Kategorien (Monat)
KR = PR1 + 5
db[f"B{KR}"] = "Wofür geht meine Zeit drauf? (dieser Monat)"; db[f"B{KR}"].font = font(size=13, bold=True, color=NAVY)
for i, h in enumerate(["Kategorie", "Stunden", "Anteil"]): header(db.cell(row=KR + 1, column=2 + i, value=h))
for i in range(15):
    rr = KR + 2 + i; k = f"Projekte!$F${5+i}"
    db[f"B{rr}"] = f'=IF({k}="","",{k})'
    db[f"C{rr}"] = f'=IF(B{rr}="","",SUMIFS({STD},{KAT},B{rr},{DAT},">="&DATE(YEAR($C$3),MONTH($C$3),1),{DAT},"<="&EOMONTH($C$3,0)))'
    db[f"D{rr}"] = f'=IF(OR(B{rr}="",$D$11=0),"",C{rr}/{cells_of["Dieser Monat"]})'
    for col in "BCD": db[f"{col}{rr}"].font = font()
    db[f"C{rr}"].number_format = "0.0;-0.0;\"–\""; db[f"D{rr}"].number_format = "0%"
    db[f"D{rr}"] = f'=IF(OR(B{rr}="",{cells_of["Dieser Monat"]}=0),"",C{rr}/{cells_of["Dieser Monat"]})'
db.conditional_formatting.add(f"D{KR+2}:D{KR+16}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="70AD47"))

# Letzte 14 Tage
TR = 48
db[f"I{TR+8}"] = "Letzte 14 Tage"; db[f"I{TR+8}"].font = font(size=13, bold=True, color=NAVY)
T0 = TR + 9
for i, h in enumerate(["Datum", "Tag", "Stunden"]): header(db.cell(row=T0, column=9 + i, value=h))
for i in range(14):
    rr = T0 + 1 + i
    db[f"I{rr}"] = f"=$C$3-{13-i}"; db[f"I{rr}"].number_format = "DD.MM."
    db[f"J{rr}"] = f'=CHOOSE(WEEKDAY(I{rr},2),"Mo","Di","Mi","Do","Fr","Sa","So")'
    db[f"K{rr}"] = f"=SUMIFS({STD},{DAT},I{rr})"; db[f"K{rr}"].number_format = "0.0"
    for col in "IJK": db[f"{col}{rr}"].font = font(); db[f"{col}{rr}"].border = Border(bottom=thin)
db.conditional_formatting.add(f"I{T0+1}:K{T0+14}", FormulaRule(formula=[f"WEEKDAY($I{T0+1},2)>5"], font=Font(color="A6A6A6")))

# Diagramme
ch = BarChart(); ch.type = "col"; ch.title = "Stunden – letzte 14 Tage"; ch.style = 10
ch.add_data(Reference(db, min_col=11, min_row=T0, max_row=T0 + 14), titles_from_data=True)
ch.set_categories(Reference(db, min_col=9, min_row=T0 + 1, max_row=T0 + 14))
ch.legend = None; ch.height = 7; ch.width = 16; ch.y_axis.title = "Std."
ch.y_axis.delete = False; ch.x_axis.delete = False; ch.x_axis.number_format = "DD.MM."
db.add_chart(ch, f"B{KR+18}")

ch2 = BarChart(); ch2.type = "bar"; ch2.title = "Stunden je Projekt (Gesamt)"; ch2.style = 10
ch2.add_data(Reference(db, min_col=5, min_row=PR0 - 1, max_row=PR0 + 9), titles_from_data=True)
ch2.set_categories(Reference(db, min_col=2, min_row=PR0, max_row=PR0 + 9))
ch2.legend = None; ch2.height = 7; ch2.width = 16
ch2.y_axis.delete = False; ch2.x_axis.delete = False; ch2.x_axis.scaling.orientation = "maxMin"
db.add_chart(ch2, f"I{T0+16}")

# Projektzeilen 16-45 – zu viele leere Zeilen gruppieren? Lieber: nur 15 sichtbar lassen
for rr in range(PR0 + 15, PR1 + 1):
    db.row_dimensions[rr].outlineLevel = 1; db.row_dimensions[rr].hidden = True
line = Border(bottom=thin)
for rng, key in [(f"B{PR0}:G{PR1}", f"$B{PR0}"), ("I16:L45", "$I16"), (f"B{KR+2}:D{KR+16}", f"$B{KR+2}")]:
    db.conditional_formatting.add(rng, FormulaRule(formula=[f'{key}<>""'], border=line))
db.page_setup.orientation = "landscape"; db.page_setup.fitToWidth = 1; db.page_setup.fitToHeight = 0
db.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.print_title_rows = "4:4"
wb.calculation.fullCalcOnLoad = True
db.freeze_panes = "A3"

# ---------------- Anleitung ----------------
an = wb.create_sheet("Anleitung")
an.column_dimensions["A"].width = 3; an.column_dimensions["B"].width = 110
an.sheet_view.showGridLines = False
lines = [
    ("Arbeitstagebuch – so geht's", "h1"),
    ("", None),
    ("1. Einmalig: Projekte anlegen", "h2"),
    ("Im Blatt 'Projekte' Ihre Projekte eintragen – optional mit Stunden-Budget und Deadline. Kategorien, Status und Prioritäten sind dort ebenfalls frei änderbar.", None),
    ("2. Täglich erfassen (ca. 10 Sekunden pro Aufgabe)", "h2"),
    ("Blatt 'Eingabe', nächste leere Zeile:  Strg + .  (Datum)  →  Tab  →  Projekt aus Liste  →  Tab  →  Kategorie  →  Tab  →  Aufgabe  →  Tab  →  Dauer, z. B. 1,5", None),
    ("Alternativ statt Dauer: Start und Ende als Uhrzeit (z. B. 8:00 und 11:15) – die Stunden werden automatisch berechnet (auch über Mitternacht).", None),
    ("Tipp: In der Auswahlliste den Anfangsbuchstaben tippen, oder mit Alt + ↓ die Liste öffnen.", None),
    ("Tipp: Strg + D kopiert die Zelle von oben (z. B. gleiches Datum oder gleiches Projekt).", None),
    ("3. Dashboard ansehen", "h2"),
    ("Zeigt Stunden heute / Woche / Monat, Wochenziel, Stunden je Projekt inkl. Budget-Verbrauch, Deadlines mit Ampel, Kategorien und die letzten 14 Tage.", None),
    ("Der Stichtag ist automatisch 'heute'. Für einen Rückblick einfach ein anderes Datum in C3 eintragen; mit =HEUTE() wieder zurücksetzen.", None),
    ("Mehr als 15 Projekte? Links neben den Zeilennummern auf [+] klicken, dann werden alle 30 Projektzeilen sichtbar.", None),
    ("", None),
    ("Farben", "h2"),
    ("Gelb = Eingabefelder   ·   Grau = wird berechnet, nicht überschreiben   ·   Grün = heutiges Datum", None),
    ("Deadline: Rot = überfällig, Orange = in den nächsten 14 Tagen   ·   Budget: Orange ab 80 %, Rot über 100 %", None),
    ("", None),
    ("Vorschläge für eine gute Zeitplanung", "h2"),
    ("• Morgens 2 Minuten: offene Aufgaben (Filter Status = 'Offen') ansehen und die 1–3 wichtigsten mit Priorität 'Hoch' markieren.", None),
    ("• Abends 2 Minuten: Tag eintragen. Lieber grob (auf 15 Min. = 0,25 gerundet) als gar nicht.", None),
    ("• Freitags: Dashboard prüfen – Wochenziel erreicht? Welches Projekt läuft über Budget? Welche Deadline kommt?", None),
    ("• Kategorie 'E-Mail / Orga' und 'Besprechung' beobachten – wenn diese über 30 % liegen, feste Blöcke dafür einplanen.", None),
    ("• Aufgaben vorab mit Status 'Offen' und ohne Dauer eintragen = einfache To-do-Liste. Erledigt? Dauer eintragen, Status ändern.", None),
    ("• Filter in der Kopfzeile von 'Eingabe' nutzen (z. B. nach Projekt oder KW), um Stundennachweise für Kunden zu erstellen.", None),
    ("", None),
    ("Die Beispielzeilen in 'Eingabe' und die Beispielprojekte können einfach gelöscht werden (Inhalte löschen mit Entf – nicht die Zeilen löschen, damit die grauen Formeln erhalten bleiben).", "note"),
]
for i, (t, s) in enumerate(lines, 1):
    c = an.cell(row=i, column=2, value=t); c.alignment = Alignment(wrap_text=True, vertical="top")
    c.font = font(size=18, bold=True, color=NAVY) if s == "h1" else font(size=12, bold=True, color=NAVY) if s == "h2" else font(italic=True, color="C00000") if s == "note" else font(size=10)

# Reiter-Farben
db.sheet_properties.tabColor = NAVY; ws.sheet_properties.tabColor = "FFC000"
ls.sheet_properties.tabColor = "70AD47"; an.sheet_properties.tabColor = "A6A6A6"
wb.active = 1  # Eingabe öffnen? -> Dashboard
wb.active = 0
import os
wb.save(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "Arbeitstagebuch.xlsx"))
print("ok")
