"""Builds the English staffing workbook (management dashboard + capacity planning input sheet).

Usage: python build_kapaplanung.py <target.xlsx>
Task values come from the original file Kapaplanung_BB2_Instandhaltung_Planug_Umsetzung_ORIGINAL.xlsm (21 Jul 2025).
"""
import sys, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, RichTextProperties
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, DataBarRule

OUT = sys.argv[1]
wb = openpyxl.Workbook()

# ---------- design ----------
NAVY, BLUE, GREY_BG, LINE, INPUT, MUTED, TOTAL = "1F3864", "2E75B6", "F3F6FA", "D9DEE7", "FFF4CC", "6B7280", "DDEBF7"
GREEN, RED = "2E7D32", "C00000"
PIE_COLORS = ["1F3864", "2E75B6", "00A3A1", "70AD47", "FFC000", "ED7D31", "8E7CC3", "7F8C9F"]
side = Side(style="thin", color=LINE)
BOX = Border(left=side, right=side, top=side, bottom=side)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
PLAIN = Alignment(vertical="center")

def font(size=10, b=False, color="000000", i=False): return Font(name="Arial", size=size, b=b, color=color, i=i)
def fill(c): return PatternFill("solid", fgColor=c)
def name(n, ref): wb.defined_names[n] = DefinedName(n, attr_text=ref)

def put(ws, ref, v, fmt=None, size=10, b=False, color="000000", bg=None, align=LEFT, border=True, i=False):
    c = ws[ref]; c.value = v; c.font = font(size, b, color, i); c.alignment = align
    if fmt: c.number_format = fmt
    if bg: c.fill = fill(bg)
    if border: c.border = BOX
    return c

def band(ws, row, c1, c2, text, size=11, bg=NAVY, height=22):
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    c = ws.cell(row, c1, text); c.font = font(size, True, "FFFFFF")
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col in range(c1, c2 + 1): ws.cell(row, col).fill = fill(bg)
    ws.row_dimensions[row].height = height

def pie(ws, cats, vals, anchor, w, h, title):
    pc = PieChart(); pc.title = title; pc.varyColors = True
    pc.add_data(vals, titles_from_data=False); pc.set_categories(cats)
    s = pc.series[0]
    for idx, col in enumerate(PIE_COLORS):
        pt = DataPoint(idx=idx); pt.graphicalProperties.solidFill = col
        pt.graphicalProperties.line.solidFill = "FFFFFF"; pt.graphicalProperties.line.width = 19050
        s.dPt.append(pt)
    s.dLbls = DataLabelList(showPercent=True, showVal=False, showCatName=False, showSerName=False, showLegendKey=False)
    s.dLbls.numFmt = "0%"
    s.dLbls.txPr = RichText(bodyPr=RichTextProperties(), p=[Paragraph(
        pPr=ParagraphProperties(defRPr=CharacterProperties(sz=1000, b=True, solidFill="FFFFFF")), endParaRPr=CharacterProperties())])
    pc.legend.position = "b"
    pc.plot_area.layout = Layout(manualLayout=ManualLayout(x=0.18, y=0.12, w=0.64, h=0.62))
    pc.width, pc.height = w, h
    ws.add_chart(pc, anchor)

# =====================================================================
# Capacity Planning (input sheet)
# =====================================================================
cp = wb.active; cp.title = "Capacity Planning"
cp.sheet_view.showGridLines = False
put(cp, "A1", "Capacity Planning – Maintenance BB2 (Implementation Team)", size=16, b=True, color=NAVY, border=False, align=PLAIN)
put(cp, "A2", "Enter only the hours per day (yellow). Week, month, year and headcount are calculated automatically.",
    color=MUTED, border=False, align=PLAIN, i=True)

heads = ["No.", "Task", "Hours\nper day", "Hours\nper week", "Hours\nper month", "Hours\nper year", "Employees\n(FTE)", "Share of\nhours", "Remarks"]
for col, h in zip("ABCDEFGHI", heads):
    put(cp, f"{col}4", h, b=True, color="FFFFFF", bg=NAVY, align=CENTER)
cp.row_dimensions[4].height = 34

tasks = [
    ("Spare parts management – planning", 1.5, None),
    ("Spare parts relocation & miscellaneous", 2, "Warehousing / external spare parts procurement"),
    ("Complex repairs / general overhauls", 0.3, None),
    ("Repair / overhaul of spare parts", 0, "Complex spare parts (pressure intensifiers, water valves …)"),
    ("Projects & continuous improvement (lean, e.g. SMED)", 10, "Development, calculation, ordering"),
    ("Training / instruction / support / TPM / WPS", 0.5, None),
    ("Assignments abroad / international support", 0.5, None),
    ("Commissioning of new equipment", 0.5, None),
    ("Approval of new tools (trials)", 0.2, None),
    ("Design / 3D CAD", 1, None),
    ("Machine relocations", 0.1, None),
    ("CNC technology (e.g. re-turning blades)", 0, None),
    ("Special tasks (storage systems …)", 0.5, None),
    ("Facility management", 6, None),
    ("Planning & documentation", 0, None),
    ("Support of bellows area", 2, None),
    ("Maintenance of operating equipment", 0, None),
    ("Maintenance of general machinery", 0, None),
    ("Administration", 3, None),
    (None, None, "Free row – add another task"),
    (None, None, "Free row – add another task"),
    (None, None, "Free row – add another task"),
]
R0 = 5; RN = R0 + len(tasks) - 1; TOT = RN + 1
for i, (task, h, rem) in enumerate(tasks):
    r = R0 + i
    bg = GREY_BG if i % 2 else None
    put(cp, f"A{r}", i + 1, "0", align=CENTER, bg=bg, color=MUTED)
    put(cp, f"B{r}", task, color="0000FF", bg=INPUT)
    put(cp, f"C{r}", h, "0.0", color="0000FF", bg=INPUT, align=CENTER, b=True)
    put(cp, f"D{r}", f'=IF(C{r}="","",C{r}*DaysPerWeek)', "#,##0.0", align=CENTER, bg=bg)
    put(cp, f"E{r}", f'=IF(C{r}="","",F{r}/MonthsPerYear)', "#,##0", align=CENTER, bg=bg)
    put(cp, f"F{r}", f'=IF(C{r}="","",C{r}*WorkDaysPerYear)', "#,##0", align=CENTER, bg=bg, b=True)
    put(cp, f"G{r}", f'=IF(C{r}="","",F{r}/HoursPerEmployee)', "0.00", align=CENTER, bg=bg)
    put(cp, f"H{r}", f'=IF(OR(C{r}="",$F${TOT}=0),"",F{r}/$F${TOT})', "0.0%", align=CENTER, bg=bg)
    put(cp, f"I{r}", rem, size=9, color=MUTED if task else "A6A6A6", bg=bg, i=not task)
    cp[f"J{r}"] = f"=N(F{r})+(1000-ROW())/1000000"  # ranking helper (hidden column)
    cp.row_dimensions[r].height = 19
for col in "ABCDEFGHI":
    put(cp, f"{col}{TOT}", None, b=True, bg=TOTAL, align=CENTER)
put(cp, f"B{TOT}", "Total", b=True, bg=TOTAL)
for col, fmt in (("C", "0.0"), ("D", "#,##0.0"), ("E", "#,##0"), ("F", "#,##0"), ("G", "0.00"), ("H", "0.0%")):
    put(cp, f"{col}{TOT}", f"=SUM({col}{R0}:{col}{RN})", fmt, b=True, bg=TOTAL, align=CENTER)
cp.row_dimensions[TOT].height = 22
cp.conditional_formatting.add(f"F{R0}:F{RN}", DataBarRule(start_type="num", start_value=0, end_type="max", color="9DC3E6"))
dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="24", allow_blank=True,
                    errorTitle="Invalid input", error="Please enter hours per day between 0 and 24.")
cp.add_data_validation(dv); dv.add(f"C{R0}:C{RN}")

# assumptions (right)
band(cp, 4, 11, 12, "Assumptions", height=34)
params = [("AsOfDate", "As of", datetime.date(2025, 7, 21), "DD MMM YYYY"),
          ("WorkDaysPerYear", "Working days per year", 250, "0"),
          ("DaysPerWeek", "Working days per week", 5, "0"),
          ("MonthsPerYear", "Months per year", 12, "0"),
          ("HoursPerEmployee", "Hours per employee per year (1 FTE)", 1750, "#,##0"),
          ("VacationDays", "Vacation days per employee", 38, "0"),
          ("RefDays", "Reference working days (vacation rate)", 220, "0"),
          ("SickRate", "Sickness rate", 0.041, "0.0%"),
          ("Headcount", "Current headcount (employees)", 4, "0.0")]
for i, (n, lab, v, fmt) in enumerate(params):
    r = 5 + i
    put(cp, f"K{r}", lab); put(cp, f"L{r}", v, fmt, color="0000FF", bg=INPUT, align=CENTER, b=True)
    name(n, f"'Capacity Planning'!$L${r}")
put(cp, "K14", "Yellow cells = input · all other cells calculate automatically", size=8, color=MUTED, border=False, align=PLAIN, i=True)

band(cp, 16, 11, 12, "Staffing result")
res = [("Employees needed for tasks", f"=G{TOT}"),
       ('="+ Vacation ("&TEXT(VacationDays/RefDays,"0.0%")&")"', f"=G{TOT}*VacationDays/RefDays"),
       ('="+ Sickness ("&TEXT(SickRate,"0.0%")&")"', f"=G{TOT}*SickRate")]
for i, (lab, v) in enumerate(res):
    r = 17 + i; put(cp, f"K{r}", lab); put(cp, f"L{r}", v, "0.00", align=CENTER)
put(cp, "K20", "Required headcount", b=True, size=11, bg=TOTAL)
put(cp, "L20", "=SUM(L17:L19)", "0.0", b=True, size=14, color=NAVY, bg=TOTAL, align=CENTER)
cp.row_dimensions[20].height = 26
put(cp, "K21", "Rounded up (persons)", b=True); put(cp, "L21", "=ROUNDUP(L20,0)", "0", b=True, align=CENTER)
put(cp, "K22", "Current headcount"); put(cp, "L22", "=Headcount", "0.0", align=CENTER)
put(cp, "K23", '=IF(L22>=L20,"Staffing reserve","Staffing gap")', b=True)
put(cp, "L23", "=L22-L20", "+0.0;-0.0;0.0", b=True, size=12, align=CENTER)
cp.conditional_formatting.add("K23:L23", FormulaRule(formula=["$L$22<$L$20"], fill=fill("F8CBAD"), font=Font(name="Arial", b=True, color="9C0006")))
cp.conditional_formatting.add("K23:L23", FormulaRule(formula=["$L$22>=$L$20"], fill=fill("C6EFCE"), font=Font(name="Arial", b=True, color="006100")))
name("TasksFTE", "'Capacity Planning'!$L$17")
name("VacationFTE", "'Capacity Planning'!$L$18")
name("SickFTE", "'Capacity Planning'!$L$19")
name("ReqHeadcount", "'Capacity Planning'!$L$20")
name("TotalHoursYear", f"'Capacity Planning'!$F${TOT}")
name("TotalHoursDay", f"'Capacity Planning'!$C${TOT}")

for col, w in zip("ABCDEFGHIJKL", (5, 46, 9, 9, 9, 10, 10, 9, 40, 3, 36, 13)):
    cp.column_dimensions[col].width = w
cp.column_dimensions["J"].hidden = True
cp.freeze_panes = "C5"

# =====================================================================
# Management Dashboard
# =====================================================================
db = wb.create_sheet("Management Dashboard", 0)
db.sheet_view.showGridLines = False
db.column_dimensions["A"].width = 2; db.column_dimensions["N"].width = 2
for col in "BCDEFGHIJKLM": db.column_dimensions[col].width = 13.5

db.merge_cells("B2:M2")
put(db, "B2", "Maintenance BB2 – Staffing & Capacity Dashboard", size=20, b=True, color="FFFFFF", bg=NAVY,
    align=Alignment(horizontal="left", vertical="center", indent=1), border=False)
for col in range(2, 14): db.cell(2, col).fill = fill(NAVY)
db.row_dimensions[2].height = 42
db.merge_cells("B3:M3")
put(db, "B3", '="Annual evaluation of the implementation team   ·   as of "&TEXT(AsOfDate,"DD MMM YYYY")&"   ·   basis: "&TEXT(WorkDaysPerYear,"0")&" working days, "&TEXT(HoursPerEmployee,"#,##0")&" h per employee and year"',
    size=10, color="FFFFFF", bg=BLUE, align=Alignment(horizontal="left", vertical="center", indent=1), border=False)
for col in range(2, 14): db.cell(3, col).fill = fill(BLUE)
db.row_dimensions[3].height = 22

def tile(row, c1, label, formula, fmt, sub, accent=NAVY):
    c2 = c1 + 2
    for rr in (row, row + 1, row + 2):
        db.merge_cells(start_row=rr, start_column=c1, end_row=rr, end_column=c2)
        for col in range(c1, c2 + 1): db.cell(rr, col).fill = fill(GREY_BG)
    for col in range(c1, c2 + 1): db.cell(row, col).border = Border(top=Side(style="thick", color=accent))
    a = db.cell(row, c1, label); a.font = font(9, True, MUTED); a.alignment = Alignment(horizontal="center", vertical="bottom")
    v = db.cell(row + 1, c1, formula); v.font = font(26, True, accent); v.alignment = CENTER; v.number_format = fmt
    s = db.cell(row + 2, c1, sub); s.font = font(8, False, MUTED, True); s.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    db.row_dimensions[row].height = 22; db.row_dimensions[row + 1].height = 40; db.row_dimensions[row + 2].height = 24

# KPI row 1 – staffing
tile(5, 2, "TOTAL HOURS PER YEAR", "=TotalHoursYear", '#,##0" h"', "all employees · all tasks")
tile(5, 5, "REQUIRED HEADCOUNT", "=ReqHeadcount", "0.0", '="incl. vacation & sickness · "&TEXT(ROUNDUP(ReqHeadcount,0),"0")&" persons rounded up"', BLUE)
tile(5, 8, "CURRENT HEADCOUNT", "=Headcount", "0.0", "employees available")
tile(5, 11, "STAFFING GAP", "=Headcount-ReqHeadcount", "+0.0;-0.0;0.0", "negative = employees missing", RED)
db.conditional_formatting.add("K6", FormulaRule(formula=["$K$6>=0"], font=Font(name="Arial", size=26, b=True, color=GREEN)))
# KPI row 2 – hours per employee
db.row_dimensions[8].height = 10
tile(9, 2, "HOURS PER EMPLOYEE / YEAR", "=HoursPerEmployee", '#,##0" h"', "contractual capacity of 1 FTE")
tile(9, 5, "AVG. WORKLOAD PER EMPLOYEE / YEAR", "=IF(Headcount=0,0,TotalHoursYear/Headcount)", '#,##0" h"', "total hours ÷ current headcount", BLUE)
tile(9, 8, "AVG. WORKLOAD PER EMPLOYEE / MONTH", "=IF(Headcount=0,0,TotalHoursYear/Headcount/MonthsPerYear)", '#,##0" h"', '="per week: "&TEXT(IF(Headcount=0,0,TotalHoursDay*DaysPerWeek/Headcount),"0.0")&" h"')
tile(9, 11, "UTILISATION", "=IF(Headcount*HoursPerEmployee=0,0,TotalHoursYear/(Headcount*HoursPerEmployee))", "0%", "task hours ÷ capacity of current staff (excl. absences)", RED)
db.conditional_formatting.add("K10", FormulaRule(formula=["$K$10<=1"], font=Font(name="Arial", size=26, b=True, color=GREEN)))

# left: hours per period
band(db, 13, 2, 7, "Hours per period")
for col, h in zip("BCDEFG", ["Period", "All employees", "Avg. per\nemployee", "Capacity\nper FTE", "Utilisation", ""]):
    put(db, f"{col}14", h, b=True, color=NAVY, bg=TOTAL, align=CENTER)
db.merge_cells("F14:G14")
db.row_dimensions[14].height = 30
periods = [("Day", "TotalHoursDay", "HoursPerEmployee/WorkDaysPerYear"),
           ("Week", "TotalHoursDay*DaysPerWeek", "HoursPerEmployee/WorkDaysPerYear*DaysPerWeek"),
           ("Month", "TotalHoursYear/MonthsPerYear", "HoursPerEmployee/MonthsPerYear"),
           ("Year", "TotalHoursYear", "HoursPerEmployee")]
for i, (p, tot, cap) in enumerate(periods):
    r = 15 + i; last = p == "Year"
    put(db, f"B{r}", p, b=True)
    put(db, f"C{r}", f"={tot}", "#,##0.0" if not last else "#,##0", align=CENTER, b=last)
    put(db, f"D{r}", f"=IF(Headcount=0,0,C{r}/Headcount)", "#,##0.0" if not last else "#,##0", align=CENTER, b=last)
    put(db, f"E{r}", f"={cap}", "#,##0.0" if not last else "#,##0", align=CENTER)
    db.merge_cells(f"F{r}:G{r}")
    put(db, f"F{r}", f"=IF(E{r}=0,0,D{r}/E{r})", "0%", align=CENTER); db[f"G{r}"].border = BOX
    db.row_dimensions[r].height = 20

# left: staffing calculation
band(db, 20, 2, 7, "Staffing calculation")
calc = [("Employees needed for tasks", "=TasksFTE", "0.00", False),
        ('="+ Vacation ("&TEXT(VacationDays/RefDays,"0.0%")&")"', "=VacationFTE", "0.00", False),
        ('="+ Sickness ("&TEXT(SickRate,"0.0%")&")"', "=SickFTE", "0.00", False),
        ("Required headcount", "=ReqHeadcount", "0.0", True),
        ("Current headcount", "=Headcount", "0.0", False),
        ('=IF(Headcount>=ReqHeadcount,"Staffing reserve","Staffing gap")', "=Headcount-ReqHeadcount", "+0.0;-0.0;0.0", True)]
for i, (lab, f, fmt, b) in enumerate(calc):
    r = 21 + i
    db.merge_cells(f"B{r}:E{r}")
    put(db, f"B{r}", lab, b=b, bg=TOTAL if lab == "Required headcount" else None)
    for col in "CDE": db[f"{col}{r}"].border = BOX
    db.merge_cells(f"F{r}:G{r}")
    put(db, f"F{r}", f, fmt, b=b, align=CENTER, size=12 if b else 10, color=NAVY if b else "000000",
        bg=TOTAL if lab == "Required headcount" else None); db[f"G{r}"].border = BOX
    db.row_dimensions[r].height = 20
db.conditional_formatting.add("B26:G26", FormulaRule(formula=["Headcount<ReqHeadcount"], fill=fill("F8CBAD"), font=Font(name="Arial", b=True, color="9C0006")))
db.conditional_formatting.add("B26:G26", FormulaRule(formula=["Headcount>=ReqHeadcount"], fill=fill("C6EFCE"), font=Font(name="Arial", b=True, color="006100")))

# left: top tasks (source of pie charts)
band(db, 28, 2, 7, "Hours by task – top 7 + others")
for col, h in zip("BCDEFG", ["Rank", "Task", "", "", "Hours / year", "Share"]):
    put(db, f"{col}29", h, b=True, color=NAVY, bg=TOTAL, align=CENTER)
db.merge_cells("C29:E29")
J = f"'Capacity Planning'!$J${R0}:$J${RN}"
for k in range(1, 8):
    r = 29 + k
    idx = f"MATCH(LARGE({J},{k}),{J},0)"
    put(db, f"B{r}", k, "0", align=CENTER, color=MUTED)
    db.merge_cells(f"C{r}:E{r}")
    put(db, f"C{r}", f"=INDEX('Capacity Planning'!$B${R0}:$B${RN},{idx})")
    for col in "DE": db[f"{col}{r}"].border = BOX
    put(db, f"F{r}", f"=N(INDEX('Capacity Planning'!$F${R0}:$F${RN},{idx}))", "#,##0", align=CENTER)
    put(db, f"G{r}", f"=IF(TotalHoursYear=0,0,F{r}/TotalHoursYear)", "0.0%", align=CENTER)
    db.row_dimensions[r].height = 19
put(db, "B37", 8, "0", align=CENTER, color=MUTED)
db.merge_cells("C37:E37"); put(db, "C37", "All other tasks")
for col in "DE": db[f"{col}37"].border = BOX
put(db, "F37", "=TotalHoursYear-SUM(F30:F36)", "#,##0", align=CENTER)
put(db, "G37", "=IF(TotalHoursYear=0,0,F37/TotalHoursYear)", "0.0%", align=CENTER)
db.merge_cells("B38:E38"); put(db, "B38", "Total", b=True, bg=TOTAL)
for col in "CDE": db[f"{col}38"].border = BOX
put(db, "F38", "=SUM(F30:F37)", "#,##0", b=True, bg=TOTAL, align=CENTER)
put(db, "G38", "=SUM(G30:G37)", "0%", b=True, bg=TOTAL, align=CENTER)
db.conditional_formatting.add("F30:F37", DataBarRule(start_type="num", start_value=0, end_type="max", color="9DC3E6"))

# right: pie chart
band(db, 13, 8, 13, "Distribution of annual hours by task")
pie(db, Reference(db, min_col=3, min_row=30, max_row=37), Reference(db, min_col=6, min_row=30, max_row=37),
    "H14", 15.6, 13.4, "Share of total hours per year")

put(db, "B40", "Enter data on sheet “Capacity Planning” (yellow cells). This dashboard updates automatically.",
    size=8, color=MUTED, border=False, align=PLAIN, i=True)

# pie chart also on the input sheet (right side)
band(cp, 25, 11, 12, "Distribution of hours")
pie(cp, Reference(db, min_col=3, min_row=30, max_row=37), Reference(db, min_col=6, min_row=30, max_row=37),
    "K26", 12.5, 11, "Share of total hours per year")

# ---------- print & finish ----------
for ws, area in ((db, "A1:N40"), (cp, f"A1:L{TOT + 1}")):
    ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = area; ws.print_options.horizontalCentered = True
    ws.oddFooter.left.text = "&F"; ws.oddFooter.center.text = "fz-pA / Mrugalla"; ws.oddFooter.right.text = "&D"
cp.print_area = "A1:M48"
db.sheet_properties.tabColor = NAVY; cp.sheet_properties.tabColor = "FFC000"
wb.active = 0
cp.sheet_view.tabSelected = False
wb.properties.title = "Maintenance BB2 – Staffing & Capacity"
wb.save(OUT)
print("ok", R0, RN, TOT)
