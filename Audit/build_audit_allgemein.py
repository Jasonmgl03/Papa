# -*- coding: utf-8 -*-
"""Erzeugt die Excel-Mappe "Audit Instandhaltung allgemein".

Aufruf:  python3 build_audit_allgemein.py
Ergebnis: Audit_IH_Allgemein_Vorlage.xlsx  (leer, zum Ausfüllen)
          Audit_IH_Allgemein_Beispiel.xlsx (mit Beispielbewertungen)

Die Prüfpunkte stehen unten in PUNKTE. Für weitere Audits (Maschinen, Anlagen)
kann dieses Skript kopiert und die Liste ausgetauscht werden.
"""
import datetime as dt
import os

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))

# ---------------------------------------------------------------- Farben
NAVY = "0F172A"
GRAU_TXT = "64748B"
GRAU_LINIE = "CBD5E1"
GRAU_HG = "F1F5F9"
INPUT_HG = "FFF7CC"
INPUT_TXT = "0000FF"
NEU_HG = "DBEAFE"
GRUEN_HG, GRUEN_TXT = "C6EFCE", "006100"
GELB_HG, GELB_TXT = "FFEB9C", "9C5700"
ROT_HG, ROT_TXT = "FFC7CE", "9C0006"
NB_HG, NB_TXT = "E2E8F0", "475569"
C_ROT, C_GELB, C_GRUEN, C_GRAU, C_HELL = "C0392B", "E0A800", "2E8B57", "94A3B8", "E2E8F0"

FONT = "Arial"


def f(size=10, bold=False, color=None, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(hex_):
    return PatternFill("solid", fgColor=hex_)


thin = Side(style="thin", color=GRAU_LINIE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
UNTEN = Border(bottom=Side(style="thin", color=NAVY))
WRAP = Alignment(wrap_text=True, vertical="top")
MITTE = Alignment(horizontal="center", vertical="center", wrap_text=True)
LINKS_M = Alignment(horizontal="left", vertical="center", wrap_text=True)

# ---------------------------------------------------------------- Inhalte
KAPITEL = [
    "Organisation & Personal",
    "Strategie & IH-System",
    "Durchführung & Schulung",
    "Auswertung & Verbesserung",
    "Dokumentation & Kommunikation",
    "Werkstatt & Betriebsmittel",
    "Material & Ersatzteile",
    "Kennzahlen",
    "Arbeits- & Umweltschutz",
]

ZIEL_VERF = "ZIEL_VERF"  # Platzhalter: Prüfpunkt-Text mit Zielwert aus Einstellungen

# (Kapitel, Check-Punkt, Prüfhinweis / Nachweise, Vorlage-Nr. oder None = NEU,
#  Check-Punkt EN, Beispielbewertung)
PUNKTE = [
    # --- Organisation & Personal
    ("Organisation & Personal", "Existiert eine Verfahrensanweisung für das Instandhaltungssystem (aktuell / wird sie gelebt)?",
     "VA vorlegen lassen: Freigabe, Revisionsstand, letzte Überprüfung. Stichprobe, ob das Vorgehen in der Praxis eingehalten wird.",
     1, "Is there a procedure for the maintenance system? (Is it currently being followed/implemented?)", "N.I.O"),
    ("Organisation & Personal", "Existiert eine Arbeitsanweisung für die vorbeugende Instandhaltung (aktuell / wird sie gelebt)?",
     "AA einsehen; an 1–2 Anlagen nachfragen, ob die Mitarbeiter sie kennen und anwenden.",
     2, "Is there a work instruction for preventive maintenance? (Is it currently being implemented/is it being followed?)", "bed. OK"),
    ("Organisation & Personal", "Existiert eine Qualifikationsmatrix / Vertretungsregelung / Struktur der IH?",
     "Organigramm, Qualifikationsmatrix (Stand?), Vertretung für Schlüsselfunktionen geregelt.",
     3, "Does a qualification matrix / substitution arrangement / structure of the maintenance department exist?", "OK"),
    ("Organisation & Personal", "Existiert eine Schichtregelung? Ist die technische Betreuung der Anlage gewährleistet?",
     "Schichtplan, IH-Besetzung je Schicht, Erreichbarkeit in allen Schichten.",
     4, "Is there a shift system in place? Is technical support for the plant guaranteed?", "n. b."),
    ("Organisation & Personal", "Führungskräfteanzahl in der IH",
     "Führungsspanne (Mitarbeiter je Führungskraft), Stellvertretung, Organigramm.",
     5, "Number of managers in maintenance", "bed. OK"),
    ("Organisation & Personal", "Existiert eine Kapazitäts- und Terminplanung der MA (T/W/M/J)?",
     "Tages-/Wochen-/Monatsplanung, Auslastung, Urlaubs- und Abwesenheitsplanung.",
     6, "Is there capacity and scheduling planning for employees (daily/weekly/monthly/yearly)?", "bed. OK"),
    ("Organisation & Personal", "Flexibilität und Qualifikation der MA innerhalb der Fertigungsbereiche",
     "Einsatz in mehreren Bereichen möglich? Mehrfachqualifikation laut Matrix.",
     7, "Flexibility and qualification of employees within the production areas", "OK"),
    ("Organisation & Personal", "Gibt es eine Rufbereitschaft und einen Eskalationsweg bei Störungen außerhalb der Regelarbeitszeit?",
     "Rufbereitschaftsplan, Eskalationsstufen, vereinbarte Reaktionszeiten, Erreichbarkeit nachts / am Wochenende.",
     None, "Is there an on-call service and escalation path for breakdowns outside regular hours?", "bed. OK"),
    ("Organisation & Personal", "Ist das Erfahrungswissen gesichert (Know-how-Träger, Nachfolge, Einarbeitung)?",
     "Altersstruktur, Know-how-Träger benannt, Einarbeitungspläne, dokumentiertes Erfahrungswissen.",
     None, "Is expert knowledge secured (key persons, succession, onboarding)?", "N.I.O"),
    ("Organisation & Personal", "Sind Elektrofachkräfte und zur Prüfung befähigte Personen schriftlich bestellt?",
     "Bestellung EFK / EuP, befähigte Personen nach BetrSichV § 2, Verantwortliche Elektrofachkraft benannt.",
     None, "Are qualified electricians and competent persons for inspections formally appointed?", "OK"),
    # --- Strategie & IH-System
    ("Strategie & IH-System", "Anwendung von Instandhaltungsstrategien (welche Instandhaltungsstrategie?)",
     "Strategie je Anlagenklasse (korrektiv, präventiv, zustandsorientiert, TPM) und Begründung.",
     8, "Application of maintenance strategies (which maintenance strategy)", "OK"),
    ("Strategie & IH-System", "Sind die Anlagen nach Kritikalität (A/B/C) eingestuft und ist die Strategie daraus abgeleitet?",
     "Anlagenliste mit A/B/C-Einstufung und Kriterien; Schlüsselanlagen nach IATF 16949, 8.5.1.5.",
     None, "Are machines classified by criticality (A/B/C) and is the strategy derived from it?", "bed. OK"),
    ("Strategie & IH-System", "Welches IH-System steht zur Verfügung (Wartung, Inspektion, geplante/ungeplante Reparatur) und wird es genutzt?",
     "System zeigen lassen (z. B. SAP PM, CMMS, Excel); Auftragsarten Wartung / Inspektion / Reparatur.",
     9, "Which maintenance system is available (maintenance; inspection; planned/unplanned repairs) and is it being used?", "bed. OK"),
    ("Strategie & IH-System", "Sind die Stammdaten im IH-System vollständig und gepflegt (Equipments, Struktur)?",
     "Stichprobe 3 Anlagen: Equipment / technischer Platz, Struktur Anlage – Baugruppe – Bauteil, Stücklisten verknüpft.",
     None, "Is master data in the maintenance system complete and maintained?", "bed. OK"),
    ("Strategie & IH-System", "Werden IH-Aufträge vollständig zurückgemeldet (Zeit, Material, Schadensbild, Ursache)?",
     "Stichprobe abgeschlossener Aufträge: Zeiten, Material, Schadensbild, Ursache und Maßnahme erfasst?",
     None, "Are work orders fully confirmed (time, material, damage, cause)?", "N.I.O"),
    ("Strategie & IH-System", "Existiert eine vorbeugende Instandhaltung?",
     "Wartungsplan je Anlage vorhanden? Anteil der Anlagen mit Plan.",
     10, "Preventive maintenance exists", "bed. OK"),
    ("Strategie & IH-System", "Qualität planmäßige Wartung (Intervalle, Arbeitspläne, Arbeitsabläufe)",
     "Intervalle begründet (Hersteller, Erfahrung)? Arbeitspläne mit Schritten, Material und Zeit.",
     11, "Quality of scheduled maintenance (intervals; work plans; work processes)", "OK"),
    ("Strategie & IH-System", "Qualität planmäßige Inspektion (Intervalle, Arbeitspläne, Arbeitsabläufe)",
     "Inspektionspläne mit Prüfkriterien und Grenzwerten, Nachweise der Durchführung.",
     12, "Quality of scheduled inspections (intervals; work plans; work procedures)", "OK"),
    ("Strategie & IH-System", "Qualität planmäßig vorbeugende Instandsetzung (Intervalle, Arbeitspläne, Arbeitsabläufe)",
     "Geplante Tauschintervalle für Verschleißteile, Generalüberholungen eingeplant.",
     13, "Quality of scheduled preventive maintenance (intervals; work plans; workflows)", "OK"),
    ("Strategie & IH-System", "Qualität störungsbedingte, bedarfsweise Instandsetzung / Einsatzbereitschaft (Arbeitspläne, Reaktionsvermögen bei Störungen)",
     "Störungsmeldeweg, Priorisierung, Reaktionszeit; Stichprobe aktueller Störmeldungen.",
     14, "Quality of breakdown-based and demand-based maintenance / operational readiness (work plans; responsiveness to malfunctions)", "OK"),
    ("Strategie & IH-System", "Werden zustandsorientierte Verfahren eingesetzt (z. B. Thermografie, Schwingungs-, Ölanalyse)?",
     "Eingesetzte Verfahren, Messrouten, Grenzwerte, Trends dokumentiert, Reaktion auf Befunde.",
     None, "Are condition-based methods used (e.g. thermography, vibration or oil analysis)?", "bed. OK"),
    ("Strategie & IH-System", "Wird Outsourcing (Fremdauftragsvergabe) für Wartung, Inspektion und geplante Instandsetzung praktiziert?",
     "Fremdvergabe-Strategie, Rahmenverträge, Leistungsbewertung der Dienstleister.",
     15, "Is outsourcing (contracting) practiced for maintenance, inspection, and planned repairs?", "OK"),
    # --- Durchführung & Schulung
    ("Durchführung & Schulung", "Sind MA für präventive Instandhaltung (PM) geschult worden?",
     "Schulungsplan und -nachweise, Unterweisungen, Herstellerschulungen.",
     16, "Have employees been trained in preventive maintenance (PM)?", "bed. OK"),
    ("Durchführung & Schulung", "Werden Betriebsstoffe für PM bereitgestellt?",
     "Schmierstoffe, Filter, Verbrauchsmaterial verfügbar und gekennzeichnet.",
     17, "Are operating materials provided for PM?", "bed. OK"),
    ("Durchführung & Schulung", "Werden Wartungen sachgemäß durchgeführt? Sichtkontrolle? Erledigte/unerledigte Aufträge?",
     "Begehung vor Ort; Anteil überfälliger Wartungsaufträge im System.",
     18, "Is maintenance being carried out properly? Visual inspection? Completed/uncompleted tasks?", "bed. OK"),
    ("Durchführung & Schulung", "Wie werden Werkzeuge gewartet? In welchem Zustand?",
     "Zustand von Hand- und Spezialwerkzeug, Prüfung, Lagerung.",
     19, "How are tools maintained? In what condition?", "OK"),
    ("Durchführung & Schulung", "Führen Bediener definierte Reinigungs-, Prüf- und Schmieraufgaben durch (autonome IH / TPM)?",
     "Bedienerstandards und Checklisten an der Anlage, Abzeichnung, Übergabe an IH bei Abweichung.",
     None, "Do operators carry out defined cleaning, inspection and lubrication tasks (autonomous maintenance / TPM)?", "bed. OK"),
    ("Durchführung & Schulung", "Gibt es nach Instandsetzung eine dokumentierte Funktionsprüfung und Freigabe vor dem Wiederanlauf?",
     "Freigabe (ggf. durch Qualität) dokumentiert, Rückmeldung an Produktion, Erstteilprüfung.",
     None, "Is there a documented function test and release before restarting after repair?", "bed. OK"),
    # --- Auswertung & Verbesserung
    ("Auswertung & Verbesserung", "Werden Wartungsmängel ausgewertet?",
     "Mängelliste aus Wartungen, Nachverfolgung, Auswertung von Häufungen.",
     20, "Are maintenance deficiencies evaluated?", "bed. OK"),
    ("Auswertung & Verbesserung", "Kontinuierliche Optimierung der Wartungs-/Inspektionspläne?",
     "Intervalle anhand von Ausfalldaten angepasst; Änderungshistorie der Pläne.",
     21, "Continuous optimization of maintenance/inspection plans?", "bed. OK"),
    ("Auswertung & Verbesserung", "Schadensregistrierung (Änderung durch konstruktives Beseitigen von Störquellen)",
     "Schäden mit Ursache erfasst; Störquellen konstruktiv beseitigt (Verbesserung nach DIN 31051).",
     22, "Damage registration (change by constructive elimination of sources of faults)", "N.I.O"),
    ("Auswertung & Verbesserung", "Schadensanalysen und Durchführung",
     "Ursachenanalysen (5-Why, Ishikawa, 8D) bei wiederkehrenden Störungen; Top-Störer-Liste.",
     23, "Damage analysis and implementation", "N.I.O"),
    ("Auswertung & Verbesserung", ZIEL_VERF,
     "Berechnung: MTBF ÷ (MTBF + MTTR). Datenquelle prüfen. Zielwert auf Blatt „Einstellungen“ (Vorlage nannte DE 90 %, EN 80 %).",
     24, "Technical availability of machines as a measure of maintenance (target see settings)", "N.I.O"),
    ("Auswertung & Verbesserung", "Gibt es regelmäßige Besprechungen IH / Produktion mit Störungs- und Maßnahmenverfolgung (Shopfloor)?",
     "Termin und Teilnehmer, Störungs- und Maßnahmenliste mit Verantwortlichen und Terminen.",
     None, "Are there regular maintenance/production meetings tracking breakdowns and actions (shopfloor)?", "OK"),
    ("Auswertung & Verbesserung", "Wird die IH bei der Beschaffung neuer Anlagen beteiligt (Lastenheft, Abnahme)?",
     "IH-Anforderungen im Lastenheft (Doku, Ersatzteile, Zugänglichkeit), Beteiligung an Abnahme.",
     None, "Is maintenance involved in purchasing new equipment (specification, acceptance)?", "bed. OK"),
    # --- Dokumentation & Kommunikation
    ("Dokumentation & Kommunikation", "Verwaltung der technischen Dokumentation (Stücklisten, Maschinenbeschreibung, technische Zeichnungen)",
     "Ablage aktuell und zugänglich (digital / Papier): Stücklisten, Schaltpläne, Zeichnungen, Betriebsanleitungen.",
     25, "Management of technical documentation (parts lists; machine descriptions; technical drawings)", "bed. OK"),
    ("Dokumentation & Kommunikation", "Wird je Anlage eine Maschinenhistorie (Lebenslaufakte) geführt?",
     "Störungen, Reparaturen, Änderungen und Prüfungen je Anlage nachvollziehbar.",
     None, "Is a machine history (log book) kept for each machine?", "bed. OK"),
    ("Dokumentation & Kommunikation", "Werden technische Änderungen an Anlagen dokumentiert (inkl. Steuerungs-/Software-Stände, Datensicherung)?",
     "Änderungsantrag und -freigabe, Unterlagen nachgeführt, aktuelle Sicherung der SPS-/Roboterprogramme.",
     None, "Are technical changes documented (incl. control/software versions, backups)?", "N.I.O"),
    ("Dokumentation & Kommunikation", "Findet eine Absprache zwischen Instandhaltung und Produktion statt?",
     "Abstimmung von Stillständen und Wartungsfenstern; Protokolle.",
     26, "Is there coordination between maintenance and production?", "OK"),
    ("Dokumentation & Kommunikation", "Wie gelangen Informationen an die Instandhalter / Anlagenführer?",
     "Schichtübergabe, Schichtbuch, Infotafel, digitale Meldungen.",
     27, "How does information reach maintenance staff / machine operators?", "OK"),
    ("Dokumentation & Kommunikation", "Wie ist die Bereitschaft der Instandhalter zu Weiterbildungsmaßnahmen?",
     "Weiterbildungsplan, Teilnahmequote, Mitarbeitergespräche.",
     28, "How willing are maintenance staff to take part in further training?", "OK"),
    # --- Werkstatt & Betriebsmittel
    ("Werkstatt & Betriebsmittel", "Wie sieht die Werkstattausstattung aus und wie ist deren Qualität?",
     "Begehung: Maschinen, Arbeitsplätze, Hebezeuge, Zustand.",
     29, "What does the workshop equipment look like and what is its quality?", "OK"),
    ("Werkstatt & Betriebsmittel", "Gibt es Betriebsmittel für Instandsetzung, PM und Generalüberholung?",
     "Vorrichtungen, Spezialwerkzeug, Prüfstände vorhanden und einsatzbereit.",
     30, "Are there resources available for repair, preventive maintenance (PM), and major overhauls?", "OK"),
    ("Werkstatt & Betriebsmittel", "Ordnung und Sauberkeit in der IH-Werkstatt (5S)?",
     "Kennzeichnung, feste Plätze (Schattenbretter), Sauberkeit; Ergebnis letztes 5S-Audit.",
     None, "Order and cleanliness in the maintenance workshop (5S)?", "bed. OK"),
    ("Werkstatt & Betriebsmittel", "Sind die Prüf- und Messmittel der IH erfasst und kalibriert?",
     "Messmittelliste (Multimeter, Drehmomentschlüssel, Messschieber …), Prüfplakette gültig.",
     None, "Are maintenance test and measuring devices registered and calibrated?", "bed. OK"),
    # --- Material & Ersatzteile
    ("Material & Ersatzteile", "Wie wird die Materialplanung (Ersatzteile, Betriebs- und Hilfsstoffe) durchgeführt (Stücklisten, Programme, Disposition)?",
     "Dispositionsverfahren, Stücklisten im System, Bestellvorschläge.",
     31, "How is material planning (spare parts, operating and auxiliary materials) carried out (bills of materials; programs; scheduling)?", "OK"),
    ("Material & Ersatzteile", "Ersatzteile Anlagen",
     "Ersatzteillisten je Anlage vorhanden, Verfügbarkeit kritischer Teile.",
     32, "Spare parts machines", "OK"),
    ("Material & Ersatzteile", "Ersatzteile Werkzeuge",
     "Ersatzteile für Werkzeuge definiert und verfügbar.",
     33, "Spare parts tools", "N.I.O"),
    ("Material & Ersatzteile", "Betriebs- und Hilfsstoffe",
     "Bevorratung, Kennzeichnung, Lagerung.",
     34, "Operating and auxiliary materials", "bed. OK"),
    ("Material & Ersatzteile", "Qualität der Ersatzteillager / Ordnung",
     "Lagerplätze gekennzeichnet, Bestand im System = Bestand im Regal (Stichprobe).",
     35, "Quality of spare parts storage / organization", "bed. OK"),
    ("Material & Ersatzteile", "Qualität der Ersatzteillager / Sicherheitsbestände",
     "Mindestbestände festgelegt und eingehalten.",
     36, "Quality of spare parts inventory / safety stocks", "bed. OK"),
    ("Material & Ersatzteile", "Sind kritische Ersatzteile für A-Anlagen festgelegt (Lieferzeit, Mindestbestand)?",
     "Liste kritischer Teile mit Lieferzeit, Mindestbestand und Lagerort; Abgleich mit Kritikalität A/B/C.",
     None, "Are critical spare parts for A-class machines defined (lead time, minimum stock)?", "N.I.O"),
    ("Material & Ersatzteile", "Werden Ersatzteile sachgerecht gelagert und konserviert (Feuchte, ESD, Haltbarkeit, FIFO)?",
     "Lagerbedingungen, Konservierung, ESD-Schutz für Elektronik, Haltbarkeit von Dichtungen/Schläuchen (IATF 16949, 8.5.1.5).",
     None, "Are spare parts stored and preserved properly (humidity, ESD, shelf life, FIFO)?", "bed. OK"),
    # --- Kennzahlen
    ("Kennzahlen", "Gibt es Instandhaltungs-Kennzahlen?",
     "z. B. MTBF, MTTR, technische Verfügbarkeit, OEE, IH-Kosten (DIN EN 15341).",
     37, "Are there maintenance indicators?", "bed. OK"),
    ("Kennzahlen", "Nutzung von Kennzahlen",
     "Werden die Kennzahlen regelmäßig besprochen und für Entscheidungen genutzt?",
     38, "Use of KPIs", "bed. OK"),
    ("Kennzahlen", "Sind Zielwerte für die IH-Kennzahlen festgelegt und wird bei Abweichung reagiert?",
     "Zielwerte dokumentiert, Soll/Ist-Vergleich, Maßnahmen bei Zielverfehlung.",
     None, "Are targets set for maintenance KPIs and is action taken on deviations?", "N.I.O"),
    ("Kennzahlen", "Sind Planungsgrad (geplante zu ungeplanter IH) und IH-Kosten (Plan/Ist) bekannt?",
     "Anteil geplanter Aufträge, IH-Budget je Kostenstelle, Plan/Ist-Vergleich.",
     None, "Are planning ratio (planned vs. unplanned) and maintenance costs (budget/actual) known?", "bed. OK"),
    # --- Arbeits- & Umweltschutz
    ("Arbeits- & Umweltschutz", "Arbeits- und Umweltschutz und technische Sicherheit",
     "Gesamteindruck: Schutzeinrichtungen, Unterweisungen, Unfallgeschehen in der IH.",
     39, "Occupational safety, environmental protection and technical safety", "bed. OK"),
    ("Arbeits- & Umweltschutz", "Sind prüfpflichtige Arbeitsmittel und Anlagen erfasst und werden die Prüffristen eingehalten?",
     "Prüfkataster (BetrSichV § 14, DGUV V3, Krane, Leitern, Druckbehälter, Tore), Nachweise, überfällige Prüfungen.",
     None, "Are equipment subject to inspection registered and inspection intervals met?", "bed. OK"),
    ("Arbeits- & Umweltschutz", "Gibt es Gefährdungsbeurteilungen für IH-Tätigkeiten und ein Freischalt-Verfahren (LOTO)?",
     "GBU nach TRBS 1112, LOTO-Verfahren, Schlösser/Anhänger vorhanden, Unterweisung nachgewiesen.",
     None, "Are there risk assessments for maintenance work and a lockout/tagout procedure?", "N.I.O"),
    ("Arbeits- & Umweltschutz", "Werden Fremdfirmen eingewiesen, koordiniert und überwacht (Erlaubnisscheine)?",
     "Fremdfirmen-Einweisung, Koordinator benannt, Erlaubnisscheine (Heißarbeiten, Befahren), Abnahme.",
     None, "Are contractors briefed, coordinated and supervised (work permits)?", "OK"),
    ("Arbeits- & Umweltschutz", "Ist der Umgang mit Gefahrstoffen und Abfällen der IH geregelt (Altöl, Emulsionen, AwSV)?",
     "Gefahrstoffverzeichnis, Sicherheitsdatenblätter, Lagerung mit Auffangwannen (WGK), Entsorgungsnachweise.",
     None, "Is handling of hazardous substances and waste regulated (waste oil, emulsions)?", "OK"),
    ("Arbeits- & Umweltschutz", "Werden Druckluft-Leckagen und Energieverluste regelmäßig gesucht und beseitigt?",
     "Leckage-Ortung (Intervall), gekennzeichnete Leckagen, Behebungsquote.",
     None, "Are compressed-air leaks and energy losses regularly detected and fixed?", "bed. OK"),
]

RESERVE = 10   # leere Zeilen für eigene Prüfpunkte
FIRST = 13     # erste Datenzeile der Checkliste
LAST = FIRST + len(PUNKTE) + RESERVE - 1

# Beispiel-Maßnahmen: Vorlage-Nr./Text-Anfang -> (Maßnahme, Verantwortlich, Tage ab heute, erledigt)
BEISPIEL_MASSN = {
    1: ("VA IH-System überarbeiten und freigeben", "Leitung IH", -20, "nein"),
    22: ("Schadenscodes im IH-System einführen", "Meister IH", 30, "nein"),
    23: ("8D-Bericht für Top-3-Störer erstellen", "Meister IH", -5, "nein"),
    24: ("Verfügbarkeit monatlich je A-Anlage auswerten", "IH-Planung", 45, "nein"),
    33: ("Ersatzteilliste Werkzeuge anlegen", "Werkzeugbau", -40, "ja"),
    2: ("Schmierplan an den Maschinen aushängen", "Meister IH", 14, "nein"),
}


def col_w(ws, widths):
    for k, v in widths.items():
        ws.column_dimensions[k].width = v


def rng(col):
    return f"Checkliste!${col}${FIRST}:${col}${LAST}"


def build(beispiel: bool, path: str):
    wb = Workbook()
    ws_d = wb.active
    ws_d.title = "Dashboard"
    ws_c = wb.create_sheet("Checkliste")
    ws_m = wb.create_sheet("Maßnahmenplan")
    ws_e = wb.create_sheet("Einstellungen")
    ws_a = wb.create_sheet("Anleitung")

    # ============================================================ Einstellungen
    ws = ws_e
    ws.sheet_view.showGridLines = False
    col_w(ws, {"A": 34, "B": 34, "C": 16, "D": 70})
    ws["A1"] = "Einstellungen"
    ws["A1"].font = f(16, True, NAVY)
    ws["A2"] = "Gelbe Felder mit blauer Schrift sind änderbar. Alle Auswertungen rechnen mit diesen Werten."
    ws["A2"].font = f(9, color=GRAU_TXT)

    ws["A4"] = "Bewertungsparameter"
    ws["A4"].font = f(11, True, NAVY)
    params = [
        (5, "Gewicht „bed. OK“", 0.5, "bed. OK zählt mit diesem Anteil (Annahme: 50 %, frei änderbar).", "Gewicht_bedOK"),
        (6, "Ampel grün ab", 0.85, "Erfüllungsgrad ab diesem Wert = „gut“ (Annahme, frei änderbar).", "Ampel_gruen"),
        (7, "Ampel gelb ab", 0.60, "Erfüllungsgrad ab diesem Wert = „mittel“, darunter „kritisch“ (Annahme).", "Ampel_gelb"),
        (8, "Ziel technische Verfügbarkeit", 0.90, "Wird im Prüfpunkt zur Verfügbarkeit angezeigt. Vorlage nannte DE 90 % / EN 80 % – bitte verbindlich festlegen.", "Ziel_Verfuegbarkeit"),
    ]
    for r, label, val, note, name in params:
        ws.cell(r, 1, label).font = f()
        c = ws.cell(r, 2, val)
        c.font = f(color=INPUT_TXT)
        c.fill = fill(INPUT_HG)
        c.number_format = "0%"
        c.border = BOX
        c.alignment = Alignment(horizontal="center")
        ws.cell(r, 4, note).font = f(9, color=GRAU_TXT)
        ws.cell(r, 4).alignment = WRAP
        wb.defined_names[name] = DefinedName(name, attr_text=f"Einstellungen!$B${r}")

    ws["A10"] = "Kapitel"
    ws["A10"].font = f(11, True, NAVY)
    ws["D10"] = ("Reihenfolge = Reihenfolge im Dashboard. Beim Umbenennen den Namen auch in der Spalte "
                 "„Kapitel“ der Checkliste anpassen.")
    ws["D10"].font = f(9, color=GRAU_TXT)
    ws["D10"].alignment = WRAP
    kap_first = 11
    for i, k in enumerate(KAPITEL):
        ws.cell(kap_first + i, 1, i + 1).font = f(color=GRAU_TXT)
        c = ws.cell(kap_first + i, 2, k)
        c.font = f(color=INPUT_TXT)
        c.fill = fill(INPUT_HG)
        c.border = BOX
    kap_last = kap_first + len(KAPITEL) - 1
    wb.defined_names["Kapitel_Liste"] = DefinedName("Kapitel_Liste", attr_text=f"Einstellungen!$B${kap_first}:$B${kap_last}")

    r0 = kap_last + 2
    ws.cell(r0, 1, "Bewertungsskala").font = f(11, True, NAVY)
    hdr = ["Bewertung", "Bedeutung", "zählt mit"]
    for j, h in enumerate(hdr):
        c = ws.cell(r0 + 1, 1 + j, h)
        c.font = f(10, True, "FFFFFF")
        c.fill = fill(NAVY)
    skala = [
        ("OK", "Anforderung erfüllt und nachweisbar", "100 %", GRUEN_HG, GRUEN_TXT),
        ("bed. OK", "Teilweise erfüllt / kleinere Abweichungen – Maßnahme nötig", "=Gewicht_bedOK", GELB_HG, GELB_TXT),
        ("N.I.O", "Nicht erfüllt / nicht vorhanden – Maßnahme nötig", "0 %", ROT_HG, ROT_TXT),
        ("n. b.", "Nicht bewertet / nicht zutreffend", "zählt nicht", NB_HG, NB_TXT),
    ]
    for i, (b, txt, z, hg, tx) in enumerate(skala):
        r = r0 + 2 + i
        c = ws.cell(r, 1, b)
        c.fill = fill(hg)
        c.font = f(10, True, tx)
        ws.cell(r, 2, txt).font = f()
        ws.cell(r, 2).alignment = WRAP
        c = ws.cell(r, 3, z)
        c.font = f()
        c.number_format = "0%"
        c.alignment = Alignment(horizontal="center")
        for j in range(1, 4):
            ws.cell(r, j).border = BOX

    r1 = r0 + 7
    ws.cell(r1, 1, "Gewichtung je Prüfpunkt").font = f(11, True, NAVY)
    for i, (g, txt) in enumerate([(1, "normal (Standard, auch wenn leer)"), (2, "wichtig"),
                                  (3, "kritisch – z. B. Sicherheit, gesetzliche Pflicht, Kundenforderung")]):
        ws.cell(r1 + 1 + i, 1, g).font = f()
        ws.cell(r1 + 1 + i, 1).alignment = Alignment(horizontal="center")
        ws.cell(r1 + 1 + i, 2, txt).font = f()
    r2 = r1 + 5
    ws.cell(r2, 1, "Auditart (Auswahlliste)").font = f(11, True, NAVY)
    arten = ["Erstaudit", "Wiederholungsaudit", "Nachaudit"]
    for i, a in enumerate(arten):
        c = ws.cell(r2 + 1 + i, 2, a)
        c.font = f(color=INPUT_TXT)
        c.fill = fill(INPUT_HG)
        c.border = BOX
    wb.defined_names["Auditart_Liste"] = DefinedName(
        "Auditart_Liste", attr_text=f"Einstellungen!$B${r2 + 1}:$B${r2 + len(arten)}")

    # ============================================================ Checkliste
    ws = ws_c
    ws.sheet_view.showGridLines = False
    col_w(ws, {"A": 6, "B": 22, "C": 50, "D": 42, "E": 10, "F": 8, "G": 11, "H": 36,
               "I": 16, "J": 11, "K": 8, "L": 11, "M": 8, "N": 8, "O": 8, "P": 8, "Q": 50})
    ws["A1"] = "Audit Instandhaltung allgemein – Checkliste"
    ws["A1"].font = f(16, True, NAVY)
    ws["A2"] = ("Gelbe Felder ausfüllen. Bewertung je Prüfpunkt aus der Liste wählen; bei N.I.O / bed. OK "
                "Befund bzw. Maßnahme, Verantwortlich und Termin eintragen. Blau hinterlegte Nummern = neue Prüfpunkte (NEU).")
    ws["A2"].font = f(9, color=GRAU_TXT)

    ws["A3"] = "Kopfdaten"
    ws["A3"].font = f(11, True, NAVY)
    kopf = [("Firma / Werk", "Beispiel GmbH, Werk 1"),
            ("Bereich / Abteilung", "Instandhaltung Fertigung"),
            ("Auditdatum", dt.date(2026, 10, 5)),
            ("Auditor", "Name Auditor"),
            ("Teilnehmer / Auditteam", "Leitung IH, Meister IH, Schichtführer"),
            ("Auditart", "Wiederholungsaudit"),
            ("Ergebnis letztes Audit", 0.55)]
    KOPF = {}
    for i, (lab, demo) in enumerate(kopf):
        r = 4 + i
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.cell(r, 1, lab).font = f()
        c = ws.cell(r, 3, demo if beispiel else None)
        c.font = f(color=INPUT_TXT)
        c.fill = fill(INPUT_HG)
        c.border = BOX
        c.alignment = LINKS_M
        KOPF[lab] = f"Checkliste!$C${r}"
    ws["C6"].number_format = "DD.MM.YYYY"
    ws["C10"].number_format = "0%"
    ws["D10"] = "Erfüllungsgrad des Vorgänger-Audits (für die Veränderung im Dashboard), leer lassen beim Erstaudit."
    ws["D10"].font = f(8, color=GRAU_TXT)
    ws["D10"].alignment = WRAP
    dv_date = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                             error="Bitte ein Datum eingeben (TT.MM.JJJJ).", showErrorMessage=True)
    dv_art = DataValidation(type="list", formula1="=Auditart_Liste", allow_blank=True)
    dv_pct = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True,
                            error="Wert zwischen 0 % und 100 % eingeben.", showErrorMessage=True)
    for dv in (dv_date, dv_art, dv_pct):
        ws.add_data_validation(dv)
    dv_date.add("C6")
    dv_art.add("C9")
    dv_pct.add("C10")

    # Live-Ergebnis rechts oben
    ws["G3"] = "Ergebnis (live)"
    ws["G3"].font = f(11, True, NAVY)
    ws.merge_cells("G4:H6")
    ws["G4"] = "=Dashboard!B8"
    ws["G4"].font = f(22, True, NAVY)
    ws["G4"].number_format = "0%"
    ws["G4"].alignment = MITTE
    ws.merge_cells("G7:H7")
    ws["G7"] = "=Dashboard!D8"
    ws["G7"].font = f(11, True)
    ws["G7"].alignment = MITTE
    ws.merge_cells("G8:H8")
    ws["G8"] = "=Dashboard!H10"
    ws["G8"].font = f(8, color=GRAU_TXT)
    ws["G8"].alignment = MITTE

    hdr = ["Nr.", "Kapitel", "Check-Punkt", "Prüfhinweis / Nachweise", "Quelle", "Gewich-\ntung",
           "Bewertung", "Befund / Maßnahme", "Verantwortlich", "Termin", "Erledigt", "Status",
           "Maßn.-Nr.", "N.I.O-Rang", "Punkte erreicht", "Punkte möglich", "Check-Punkt (EN)"]
    HR = FIRST - 1
    for j, h in enumerate(hdr, start=1):
        c = ws.cell(HR, j, h)
        c.font = f(10, True, "FFFFFF")
        c.fill = fill(NAVY if j <= 13 else "475569")
        c.alignment = MITTE
        c.border = BOX
    ws.row_dimensions[HR].height = 30
    ws.cell(HR - 1, 14, "Hilfswerte (nicht ändern)").font = f(8, color=GRAU_TXT, italic=True)

    heute = dt.date(2026, 10, 5)
    for i in range(len(PUNKTE) + RESERVE):
        r = FIRST + i
        G, F_ = f"G{r}", f"F{r}"
        if i < len(PUNKTE):
            kap, txt, hint, vnr, en, demo = PUNKTE[i]
            if txt == ZIEL_VERF:
                txt = ('="Technische Verfügbarkeit der Maschinen als Maß für die IH ≥ "'
                       '&TEXT(Ziel_Verfuegbarkeit,"0%")&" (Zielwert, siehe Einstellungen)"')
            ws.cell(r, 2, kap)
            ws.cell(r, 3, txt)
            ws.cell(r, 4, hint)
            ws.cell(r, 5, f"Vorlage {vnr}" if vnr else "NEU")
            ws.cell(r, 17, en)
            if beispiel:
                ws.cell(r, 6, 2 if kap == "Arbeits- & Umweltschutz" else 1)
                ws.cell(r, 7, demo)
                if vnr in BEISPIEL_MASSN:
                    m, v, d, e = BEISPIEL_MASSN[vnr]
                    ws.cell(r, 8, m)
                    ws.cell(r, 9, v)
                    ws.cell(r, 10, heute + dt.timedelta(days=d))
                    ws.cell(r, 11, e)
            else:
                ws.cell(r, 6, 1)
        ws.cell(r, 1, f'=IF(C{r}="","",COUNTIF($C${FIRST}:C{r},"?*"))')
        ws.cell(r, 12, f'=IF(OR({G}="N.I.O",{G}="bed. OK"),IF(K{r}="ja","erledigt",'
                       f'IF(AND(J{r}<>"",J{r}<TODAY()),"überfällig","offen")),"")')
        ws.cell(r, 13, f'=IF(OR({G}="N.I.O",{G}="bed. OK"),COUNTIF($G${FIRST}:{G},"N.I.O")'
                       f'+COUNTIF($G${FIRST}:{G},"bed. OK"),"")')
        ws.cell(r, 14, f'=IF({G}="N.I.O",COUNTIF($G${FIRST}:{G},"N.I.O"),"")')
        w = f'IF({F_}="",1,{F_})'
        ws.cell(r, 15, f'=IF({G}="OK",{w},IF({G}="bed. OK",{w}*Gewicht_bedOK,IF({G}="N.I.O",0,"")))')
        ws.cell(r, 16, f'=IF(OR({G}="OK",{G}="bed. OK",{G}="N.I.O"),{w},"")')

        # Format
        for j in range(1, 18):
            c = ws.cell(r, j)
            c.border = BOX
            c.alignment = WRAP
            c.font = f(9 if j in (4, 17) else 10, color=GRAU_TXT if j in (4, 14, 15, 16, 17) else None)
        for j in (2, 3, 6, 7, 8, 9, 10, 11):   # Eingaben
            ws.cell(r, j).fill = fill(INPUT_HG)
            ws.cell(r, j).font = f(10, color=INPUT_TXT)
        if i >= len(PUNKTE):
            ws.cell(r, 4).fill = fill(INPUT_HG)
            ws.cell(r, 4).font = f(9, color=INPUT_TXT)
        ws.cell(r, 1).alignment = Alignment(horizontal="center", vertical="top")
        for j in (5, 6, 7, 10, 11, 12, 13, 14, 15, 16):
            ws.cell(r, j).alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        ws.cell(r, 10).number_format = "DD.MM.YYYY"
        ws.cell(r, 15).number_format = "0.0"
        if i < len(PUNKTE) and PUNKTE[i][3] is None:
            ws.cell(r, 1).fill = fill(NEU_HG)
            ws.cell(r, 5).fill = fill(NEU_HG)
            ws.cell(r, 5).font = f(9, True, "1D4ED8")
        else:
            ws.cell(r, 5).font = f(9, color=GRAU_TXT)

    dv_bew = DataValidation(type="list", formula1='"N.I.O,bed. OK,OK,n. b."', allow_blank=True,
                            error="Bitte N.I.O, bed. OK, OK oder n. b. wählen.", showErrorMessage=True)
    dv_ja = DataValidation(type="list", formula1='"ja,nein"', allow_blank=True)
    dv_gew = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True)
    dv_kap = DataValidation(type="list", formula1="=Kapitel_Liste", allow_blank=True)
    dv_t = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                          error="Bitte ein Datum eingeben (TT.MM.JJJJ).", showErrorMessage=True)
    for dv in (dv_bew, dv_ja, dv_gew, dv_kap, dv_t):
        ws.add_data_validation(dv)
    dv_bew.add(f"G{FIRST}:G{LAST}")
    dv_ja.add(f"K{FIRST}:K{LAST}")
    dv_gew.add(f"F{FIRST}:F{LAST}")
    dv_kap.add(f"B{FIRST}:B{LAST}")
    dv_t.add(f"J{FIRST}:J{LAST}")

    def bew_cf(ws, area):
        for val, hg, tx in (("N.I.O", ROT_HG, ROT_TXT), ("bed. OK", GELB_HG, GELB_TXT),
                            ("OK", GRUEN_HG, GRUEN_TXT), ("n. b.", NB_HG, NB_TXT)):
            ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=[f'"{val}"'],
                                                           fill=fill(hg), font=Font(name=FONT, bold=True, color=tx)))

    def status_cf(ws, area):
        for val, hg, tx in (("überfällig", ROT_HG, ROT_TXT), ("offen", GELB_HG, GELB_TXT),
                            ("erledigt", GRUEN_HG, GRUEN_TXT)):
            ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=[f'"{val}"'],
                                                           fill=fill(hg), font=Font(name=FONT, color=tx)))

    bew_cf(ws, f"G{FIRST}:G{LAST}")
    status_cf(ws, f"L{FIRST}:L{LAST}")
    ws.freeze_panes = f"D{FIRST}"
    ws.auto_filter.ref = f"A{HR}:Q{LAST}"
    ws.column_dimensions.group("N", "P", hidden=True, outline_level=1)
    ws.column_dimensions.group("Q", "Q", hidden=True, outline_level=1)
    ws.cell(LAST + 2, 1, f"Zeilen {FIRST + len(PUNKTE)}–{LAST} sind Reserve für eigene Prüfpunkte "
                         "(Kapitel wählen, Text eintragen). Nicht benötigte Prüfpunkte: ganze Zeile löschen – "
                         "Nummern und Auswertung passen sich an.").font = f(9, color=GRAU_TXT, italic=True)
    ws.cell(LAST + 3, 1, "Ausgeblendete Spalten N–Q (Hilfswerte, englischer Text): über das „+“ oberhalb "
                         "der Spaltenköpfe einblenden.").font = f(9, color=GRAU_TXT, italic=True)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{HR}:{HR}"
    ws.print_area = f"A1:M{LAST}"
    ws.oddFooter.center.text = "Seite &P von &N"

    # ============================================================ Dashboard
    ws = ws_d
    ws.sheet_view.showGridLines = False
    col_w(ws, {"A": 2, "B": 9, **{c: 12.5 for c in "CDEFGHIJKLM"}, "N": 2})
    ws["B1"] = "Audit Instandhaltung allgemein – Bewertung"
    ws["B1"].font = f(18, True, NAVY)
    ws["B2"] = "Rechnet automatisch aus dem Blatt „Checkliste“. Grenzwerte und Gewichte auf „Einstellungen“."
    ws["B2"].font = f(9, color=GRAU_TXT)

    # Kopf-Leiste
    infos = [("B", "C", "Firma / Werk", KOPF["Firma / Werk"]),
             ("D", "E", "Bereich / Abteilung", KOPF["Bereich / Abteilung"]),
             ("F", "G", "Auditdatum", KOPF["Auditdatum"]),
             ("H", "I", "Auditor", KOPF["Auditor"]),
             ("J", "K", "Auditart", KOPF["Auditart"]),
             ("L", "M", "Teilnehmer", KOPF["Teilnehmer / Auditteam"])]
    for a, b, lab, ref in infos:
        ws.merge_cells(f"{a}4:{b}4")
        ws.merge_cells(f"{a}5:{b}5")
        ws[f"{a}4"] = lab
        ws[f"{a}4"].font = f(8, color=GRAU_TXT)
        ws[f"{a}5"] = f'=IF({ref}="","–",{ref})'
        ws[f"{a}5"].font = f(10, True, NAVY)
        ws[f"{a}5"].alignment = Alignment(vertical="top", wrap_text=True)
        for c in (a, b):
            ws[f"{c}4"].fill = fill(GRAU_HG)
            ws[f"{c}5"].fill = fill(GRAU_HG)
    ws["F5"].number_format = "DD.MM.YYYY"
    ws["F5"].alignment = Alignment(horizontal="left", vertical="top")
    ws.row_dimensions[5].height = 28

    # KPI-Kacheln (Zeile 7 Titel, 8-9 Wert, 10 Untertitel)
    O, P, G_, L_ = rng("O"), rng("P"), rng("G"), rng("L")
    cnt = lambda v: f'COUNTIF({G_},"{v}")'
    bewertet = f'({cnt("N.I.O")}+{cnt("bed. OK")}+{cnt("OK")}+{cnt("n. b.")})'
    gesamt = f'COUNTIF({rng("C")},"?*")'
    prev = KOPF["Ergebnis letztes Audit"]
    tiles = [
        ("B", "C", "Erfüllungsgrad", f"=IFERROR(SUM({O})/SUM({P}),\"\")", "0%",
         '="Ziel „gut“ ≥ "&TEXT(Ampel_gruen,"0%")'),
        ("D", "E", "Gesamtbewertung",
         '=IF(B8="","nicht bewertet",IF(B8>=Ampel_gruen,"gut",IF(B8>=Ampel_gelb,"mittel","kritisch")))', None,
         '="gelb ab "&TEXT(Ampel_gelb,"0%")&" · grün ab "&TEXT(Ampel_gruen,"0%")'),
        ("F", "G", "Veränderung zum letzten Audit",
         f'=IF(OR({prev}="",B8=""),"–",(B8-{prev})*100)', '+0" Pkt.";-0" Pkt.";"±0 Pkt."',
         f'="letztes Audit: "&IF({prev}="","–",TEXT({prev},"0%"))'),
        ("H", "I", "Bewertungsfortschritt", f"=IFERROR({bewertet}/{gesamt},\"\")", "0%",
         f'={bewertet}&" von "&{gesamt}&" Punkten bewertet"'),
        ("J", "K", "Maßnahmen gesamt",
         f'=COUNTIF({L_},"offen")+COUNTIF({L_},"überfällig")+COUNTIF({L_},"erledigt")', "0",
         f'="davon erledigt: "&COUNTIF({L_},"erledigt")'),
        ("L", "M", "Maßnahmen überfällig", f'=COUNTIF({L_},"überfällig")', "0",
         f'="offen (im Termin / ohne Termin): "&COUNTIF({L_},"offen")'),
    ]
    for a, b, title, formula, fmt, sub in tiles:
        ws.merge_cells(f"{a}7:{b}7")
        ws.merge_cells(f"{a}8:{b}9")
        ws.merge_cells(f"{a}10:{b}10")
        ws[f"{a}7"] = title
        ws[f"{a}7"].font = f(9, True, GRAU_TXT)
        ws[f"{a}7"].alignment = MITTE
        ws[f"{a}8"] = formula
        ws[f"{a}8"].font = f(24 if a != "D" else 18, True, NAVY)
        ws[f"{a}8"].alignment = MITTE
        if fmt:
            ws[f"{a}8"].number_format = fmt
        ws[f"{a}10"] = sub
        ws[f"{a}10"].font = f(8, color=GRAU_TXT)
        ws[f"{a}10"].alignment = MITTE
        for rr in range(7, 11):
            for cc in (a, b):
                cell = ws[f"{cc}{rr}"]
                cell.border = Border(left=thin if cc == a else None, right=thin if cc == b else None,
                                     top=thin if rr == 7 else None, bottom=thin if rr == 10 else None)
    ws.row_dimensions[8].height = 22
    ws.row_dimensions[9].height = 22
    # Ampel-Farben Kacheln
    for area, ref in (("B8:C9", "$B$8"), ("D8:E9", "$B$8")):
        ws.conditional_formatting.add(area, FormulaRule(formula=[f'AND({ref}<>"",{ref}>=Ampel_gruen)'],
                                                        fill=fill(GRUEN_HG), font=Font(name=FONT, color=GRUEN_TXT, bold=True)))
        ws.conditional_formatting.add(area, FormulaRule(formula=[f'AND({ref}<>"",{ref}>=Ampel_gelb,{ref}<Ampel_gruen)'],
                                                        fill=fill(GELB_HG), font=Font(name=FONT, color=GELB_TXT, bold=True)))
        ws.conditional_formatting.add(area, FormulaRule(formula=[f'AND({ref}<>"",{ref}<Ampel_gelb)'],
                                                        fill=fill(ROT_HG), font=Font(name=FONT, color=ROT_TXT, bold=True)))
    ws.conditional_formatting.add("F8:G9", FormulaRule(formula=['AND(ISNUMBER($F$8),$F$8>0)'],
                                                       font=Font(name=FONT, color=GRUEN_TXT, bold=True)))
    ws.conditional_formatting.add("F8:G9", FormulaRule(formula=['AND(ISNUMBER($F$8),$F$8<0)'],
                                                       font=Font(name=FONT, color=ROT_TXT, bold=True)))
    ws.conditional_formatting.add("L8:M9", FormulaRule(formula=['$L$8>0'], fill=fill(ROT_HG),
                                                       font=Font(name=FONT, color=ROT_TXT, bold=True)))

    # Kapitel-Tabelle
    def section(row, text):
        ws.merge_cells(f"B{row}:M{row}")
        ws[f"B{row}"] = text
        ws[f"B{row}"].font = f(12, True, NAVY)
        for c in "BCDEFGHIJKLM":
            ws[f"{c}{row}"].border = UNTEN

    section(12, "Ergebnis je Kapitel")
    TH = 13
    heads = [("B", "D", "Kapitel"), ("E", "E", "Punkte"), ("F", "F", "N.I.O"), ("G", "G", "bed. OK"),
             ("H", "H", "OK"), ("I", "I", "n. b."), ("J", "J", "offen"), ("K", "K", "Erfüllungsgrad"),
             ("L", "M", "Bewertung")]
    for a, b, h in heads:
        if a != b:
            ws.merge_cells(f"{a}{TH}:{b}{TH}")
        ws[f"{a}{TH}"] = h
        ws[f"{a}{TH}"].font = f(9, True, "FFFFFF")
        ws[f"{a}{TH}"].alignment = MITTE
        for cc in range(ord(a), ord(b) + 1):
            ws[f"{chr(cc)}{TH}"].fill = fill(NAVY)
    K0 = TH + 1
    B_, C_ = rng("B"), rng("C")
    for i in range(len(KAPITEL)):
        r = K0 + i
        er = kap_first + i
        ws.merge_cells(f"B{r}:D{r}")
        ws.merge_cells(f"L{r}:M{r}")
        ws[f"B{r}"] = f'=IF(Einstellungen!$B${er}="","",Einstellungen!$B${er})'
        ws[f"E{r}"] = f'=IF($B{r}="","",COUNTIFS({B_},$B{r},{C_},"?*"))'
        for col, v in (("F", "N.I.O"), ("G", "bed. OK"), ("H", "OK"), ("I", "n. b.")):
            ws[f"{col}{r}"] = f'=IF($B{r}="","",COUNTIFS({B_},$B{r},{G_},"{v}"))'
        ws[f"J{r}"] = f'=IF($B{r}="","",E{r}-SUM(F{r}:I{r}))'
        ws[f"K{r}"] = f'=IF($B{r}="","",IFERROR(SUMIFS({O},{B_},$B{r})/SUMIFS({P},{B_},$B{r}),""))'
        ws[f"L{r}"] = (f'=IF($B{r}="","",IF(K{r}="","nicht bewertet",IF(K{r}>=Ampel_gruen,"gut",'
                       f'IF(K{r}>=Ampel_gelb,"mittel","kritisch"))))')
        for c in "BCDEFGHIJKLM":
            cell = ws[f"{c}{r}"]
            cell.font = f(10)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(horizontal="left" if c == "B" else "center", vertical="center")
            if i % 2:
                cell.fill = fill("F8FAFC")
        ws[f"K{r}"].number_format = "0%"
    KL = K0 + len(KAPITEL) - 1
    TR = KL + 1
    ws.merge_cells(f"B{TR}:D{TR}")
    ws.merge_cells(f"L{TR}:M{TR}")
    ws[f"B{TR}"] = "Gesamt"
    for c in "EFGHIJ":
        ws[f"{c}{TR}"] = f"=SUM({c}{K0}:{c}{KL})"
    ws[f"K{TR}"] = "=B8"
    ws[f"L{TR}"] = "=D8"
    for c in "BCDEFGHIJKLM":
        cell = ws[f"{c}{TR}"]
        cell.font = f(10, True, NAVY)
        cell.fill = fill(GRAU_HG)
        cell.border = Border(top=Side(style="thin", color=NAVY), bottom=Side(style="thin", color=NAVY))
        cell.alignment = Alignment(horizontal="left" if c == "B" else "center", vertical="center")
    ws[f"K{TR}"].number_format = "0%"
    ws.conditional_formatting.add(f"K{K0}:K{KL}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                             end_value=1, color="94A3B8"))
    amp = f"L{K0}:M{TR}"
    for val, hg, tx in (("gut", GRUEN_HG, GRUEN_TXT), ("mittel", GELB_HG, GELB_TXT), ("kritisch", ROT_HG, ROT_TXT)):
        ws.conditional_formatting.add(amp, CellIsRule(operator="equal", formula=[f'"{val}"'], fill=fill(hg),
                                                      font=Font(name=FONT, bold=True, color=tx)))
    for c, hg in (("F", ROT_HG), ("G", GELB_HG)):
        ws.conditional_formatting.add(f"{c}{K0}:{c}{KL}", CellIsRule(operator="greaterThan", formula=["0"],
                                                                     font=Font(name=FONT, bold=True,
                                                                               color=ROT_TXT if c == "F" else GELB_TXT)))

    # Diagramme
    CH = TR + 2
    section(CH, "Diagramme")
    cats = Reference(ws, min_col=2, min_row=K0, max_row=KL)

    ch1 = BarChart()
    ch1.type = "bar"
    ch1.style = 10
    ch1.title = "Erfüllungsgrad je Kapitel"
    ch1.add_data(Reference(ws, min_col=11, min_row=K0, max_row=KL), titles_from_data=False)
    ch1.set_categories(cats)
    ch1.legend = None
    ch1.y_axis.scaling.min = 0
    ch1.y_axis.scaling.max = 1
    ch1.y_axis.majorUnit = 0.25
    ch1.y_axis.number_format = "0%"
    ch1.y_axis.majorGridlines = None
    ch1.x_axis.scaling.orientation = "maxMin"   # Reihenfolge wie Tabelle
    ch1.x_axis.delete = False
    ch1.y_axis.delete = False
    s = ch1.series[0]
    s.graphicalProperties = GraphicalProperties(solidFill="334155")
    s.graphicalProperties.line.solidFill = "334155"
    s.dLbls = DataLabelList()
    s.dLbls.showVal = True
    for attr in ("showSerName", "showCatName", "showLegendKey", "showPercent", "showLeaderLines"):
        setattr(s.dLbls, attr, False)
    s.dLbls.position = "outEnd"
    s.dLbls.numFmt = "0%"
    ch1.gapWidth = 60
    ch1.width = 13.8
    ch1.height = 9.0
    ws.add_chart(ch1, f"B{CH + 2}")

    ch2 = BarChart()
    ch2.type = "bar"
    ch2.grouping = "stacked"
    ch2.overlap = 100
    ch2.title = "Bewertungen je Kapitel (Anzahl)"
    for col, color in ((6, C_ROT), (7, C_GELB), (8, C_GRUEN), (9, C_GRAU), (10, C_HELL)):
        ch2.add_data(Reference(ws, min_col=col, min_row=TH, max_row=KL), titles_from_data=True)
        sr = ch2.series[-1]
        sr.graphicalProperties = GraphicalProperties(solidFill=color)
        sr.graphicalProperties.line.solidFill = "FFFFFF"
    ch2.set_categories(cats)
    ch2.x_axis.scaling.orientation = "maxMin"
    ch2.x_axis.delete = False
    ch2.y_axis.delete = False
    ch2.y_axis.majorGridlines = None
    ch2.legend.position = "b"
    ch2.gapWidth = 60
    ch2.width = 13.8
    ch2.height = 9.0
    ws.add_chart(ch2, f"H{CH + 2}")

    # Kritische Befunde
    FB = CH + 20
    section(FB, "Kritische Befunde (N.I.O) – vorrangig bearbeiten")
    FH = FB + 1
    fheads = [("B", "B", "Nr."), ("C", "G", "Check-Punkt"), ("H", "J", "Befund / Maßnahme"),
              ("K", "K", "Verantwortlich"), ("L", "L", "Termin"), ("M", "M", "Status")]
    for a, b, h in fheads:
        if a != b:
            ws.merge_cells(f"{a}{FH}:{b}{FH}")
        ws[f"{a}{FH}"] = h
        ws[f"{a}{FH}"].font = f(9, True, "FFFFFF")
        ws[f"{a}{FH}"].alignment = MITTE
        for cc in range(ord(a), ord(b) + 1):
            ws[f"{chr(cc)}{FH}"].fill = fill(NAVY)
    NMAX = 15
    N_ = rng("N")
    for k in range(1, NMAX + 1):
        r = FH + k
        ws.merge_cells(f"C{r}:G{r}")
        ws.merge_cells(f"H{r}:J{r}")
        m = f"MATCH({k},{N_},0)"
        ws[f"B{r}"] = f'=IFERROR(INDEX({rng("A")},{m}),"")'
        ws[f"C{r}"] = f'=IFERROR(INDEX({rng("C")},{m}),"")'
        ws[f"H{r}"] = f'=IFERROR(INDEX({rng("H")},{m})&"","")'
        ws[f"K{r}"] = f'=IFERROR(INDEX({rng("I")},{m})&"","")'
        ws[f"L{r}"] = f'=IFERROR(IF(INDEX({rng("J")},{m})="","",INDEX({rng("J")},{m})),"")'
        ws[f"M{r}"] = f'=IFERROR(INDEX({rng("L")},{m}),"")'
        for c in "BCDEFGHIJKLM":
            cell = ws[f"{c}{r}"]
            cell.font = f(9)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(horizontal="center" if c in "BLM" else "left", vertical="top",
                                       wrap_text=True)
        ws[f"L{r}"].number_format = "DD.MM.YYYY"
        ws.row_dimensions[r].height = 26
    FL = FH + NMAX
    status_cf(ws, f"M{FH + 1}:M{FL}")
    ws.merge_cells(f"B{FL + 1}:M{FL + 1}")
    ws[f"B{FL + 1}"] = (f'=IF({cnt("N.I.O")}=0,"Keine N.I.O-Befunde.",IF({cnt("N.I.O")}>{NMAX},'
                        f'"+ "&({cnt("N.I.O")}-{NMAX})&" weitere N.I.O-Befunde – siehe Blatt „Maßnahmenplan“.",'
                        f'"Alle Maßnahmen (auch bed. OK) im Blatt „Maßnahmenplan“."))')
    ws[f"B{FL + 1}"].font = f(9, color=GRAU_TXT, italic=True)

    # Fazit & Unterschriften
    FZ = FL + 3
    section(FZ, "Fazit des Auditors")
    ws.merge_cells(f"B{FZ + 1}:M{FZ + 5}")
    c = ws[f"B{FZ + 1}"]
    c.fill = fill(INPUT_HG)
    c.font = f(10, color=INPUT_TXT)
    c.alignment = WRAP
    if beispiel:
        c.value = ("Grundstrukturen der IH sind vorhanden, die vorbeugende Instandhaltung wird aber nicht durchgängig "
                   "gelebt. Schwerpunkte: Schadensanalyse, Kennzahlen mit Zielwerten und Freischalt-Verfahren (LOTO). "
                   "Nachaudit in 6 Monaten empfohlen.")
    for rr in range(FZ + 1, FZ + 6):
        for cc in "BCDEFGHIJKLM":
            ws[f"{cc}{rr}"].fill = fill(INPUT_HG)
    SG = FZ + 9
    for a, b, txt in (("B", "F", "Datum, Unterschrift Auditor"), ("H", "M", "Datum, Unterschrift Bereichsleitung")):
        ws.merge_cells(f"{a}{SG}:{b}{SG}")
        ws[f"{a}{SG}"] = txt
        ws[f"{a}{SG}"].font = f(8, color=GRAU_TXT)
        for cc in range(ord(a), ord(b) + 1):
            ws[f"{chr(cc)}{SG}"].border = Border(top=Side(style="thin", color=NAVY))

    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.print_area = f"A1:N{SG}"
    ws.row_breaks.append(__import__("openpyxl").worksheet.pagebreak.Break(id=FB - 1))
    ws.oddFooter.center.text = "Seite &P von &N"

    # ============================================================ Maßnahmenplan
    ws = ws_m
    ws.sheet_view.showGridLines = False
    col_w(ws, {"A": 8, "B": 6, "C": 22, "D": 48, "E": 10, "F": 40, "G": 16, "H": 11, "I": 11, "J": 6})
    ws["A1"] = "Maßnahmenplan – Audit Instandhaltung allgemein"
    ws["A1"].font = f(16, True, NAVY)
    ws["A2"] = ("Wird automatisch aus der Checkliste erzeugt (alle N.I.O und bed. OK). "
                "Änderungen bitte in der Checkliste eintragen.")
    ws["A2"].font = f(9, color=GRAU_TXT)
    mh = ["Maßn.-Nr.", "Nr.", "Kapitel", "Check-Punkt", "Bewertung", "Befund / Maßnahme",
          "Verantwortlich", "Termin", "Status", "Zeile"]
    for j, h in enumerate(mh, start=1):
        c = ws.cell(4, j, h)
        c.font = f(10, True, "FFFFFF")
        c.fill = fill(NAVY if j < 10 else "475569")
        c.alignment = MITTE
        c.border = BOX
    MMAX = len(PUNKTE) + RESERVE
    for k in range(1, MMAX + 1):
        r = 4 + k
        ws.cell(r, 10, f'=IFERROR(MATCH({k},{rng("M")},0),"")')
        ws.cell(r, 1, f'=IF($J{r}="","",{k})')
        ws.cell(r, 2, f'=IF($J{r}="","",INDEX({rng("A")},$J{r}))')
        for j, src in ((3, "B"), (4, "C"), (5, "G"), (6, "H"), (7, "I"), (9, "L")):
            ws.cell(r, j, f'=IF($J{r}="","",INDEX({rng(src)},$J{r})&"")')
        ws.cell(r, 8, f'=IF($J{r}="","",IF(INDEX({rng("J")},$J{r})="","",INDEX({rng("J")},$J{r})))')
        for j in range(1, 11):
            c = ws.cell(r, j)
            c.font = f(9, color=GRAU_TXT if j == 10 else None)
            c.border = Border(bottom=thin)
            c.alignment = Alignment(horizontal="center" if j in (1, 2, 5, 8, 9, 10) else "left",
                                    vertical="top", wrap_text=True)
        ws.cell(r, 8).number_format = "DD.MM.YYYY"
    bew_cf(ws, f"E5:E{4 + MMAX}")
    status_cf(ws, f"I5:I{4 + MMAX}")
    ws.column_dimensions.group("J", "J", hidden=True, outline_level=1)
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:4"
    ws.print_area = f"A1:I{4 + MMAX}"

    # ============================================================ Anleitung
    ws = ws_a
    ws.sheet_view.showGridLines = False
    col_w(ws, {"A": 26, "B": 110})
    ws["A1"] = "Anleitung"
    ws["A1"].font = f(16, True, NAVY)
    rows = [
        ("So geht's", None),
        ("1.", "Blatt „Checkliste“: Kopfdaten ausfüllen (Firma, Bereich, Datum, Auditor, Teilnehmer, Auditart, Ergebnis letztes Audit)."),
        ("2.", "Je Prüfpunkt in „Bewertung“ wählen: OK, bed. OK, N.I.O oder n. b. (nicht bewertet / nicht zutreffend). "
               "Die Spalte „Prüfhinweis / Nachweise“ sagt, worauf zu achten ist."),
        ("3.", "Bei N.I.O und bed. OK: Befund / Maßnahme, Verantwortlich und Termin eintragen. Erledigt = ja setzt den Status auf „erledigt“; "
               "Termin in der Vergangenheit = „überfällig“."),
        ("4.", "Optional: Gewichtung 1–3 je Prüfpunkt (leer = 1). Wichtige Punkte zählen dann stärker."),
        ("5.", "Blatt „Dashboard“ zeigt das Ergebnis: Erfüllungsgrad, Ampel, Kapitel, Diagramme, kritische Befunde. "
               "Fazit eintragen, drucken, unterschreiben."),
        ("6.", "Blatt „Maßnahmenplan“ listet alle Maßnahmen automatisch – zum Nachverfolgen oder Weitergeben."),
        (None, None),
        ("Farben", None),
        ("Gelb, blaue Schrift", "Eingabefelder – nur hier eintragen."),
        ("Schwarz / grau", "Formeln – nicht überschreiben."),
        ("Blaue Nummer, „NEU“", "Zusätzlich vorgeschlagene Prüfpunkte (nicht in der ursprünglichen Vorlage). "
                                "Nicht benötigte Punkte: ganze Zeile löschen (Rechtsklick auf Zeilennummer → Zellen löschen)."),
        (None, None),
        ("Berechnung", None),
        ("Erfüllungsgrad", "Σ erreichte Punkte ÷ Σ mögliche Punkte. OK = Gewichtung × 100 %, bed. OK = Gewichtung × Gewicht „bed. OK“ "
                           "(Annahme 50 %), N.I.O = 0. „n. b.“ und unbewertete Punkte zählen nicht."),
        ("Ampel", "gut ≥ 85 %, mittel ≥ 60 %, sonst kritisch. Annahmen – auf „Einstellungen“ änderbar."),
        ("Veränderung", "Erfüllungsgrad jetzt minus „Ergebnis letztes Audit“ in Prozentpunkten."),
        ("Fortschritt", "Anteil der Prüfpunkte mit einer Bewertung (inkl. n. b.)."),
        (None, None),
        ("Eigene Prüfpunkte", "Am Ende der Checkliste sind 10 leere Reservezeilen. Kapitel wählen und Text eintragen – Auswertung "
                              "und Dashboard zählen sie automatisch mit. Weitere Zeilen innerhalb des Bereichs einfügen "
                              "(nicht unterhalb der letzten Reservezeile)."),
        ("Neues Kapitel", "Auf „Einstellungen“ einen Kapitelnamen ersetzen. Das Dashboard zeigt 9 Kapitel."),
        ("Nächstes Audit", "Datei kopieren (Speichern unter …), Bewertungen leeren, Erfüllungsgrad des alten Audits bei "
                           "„Ergebnis letztes Audit“ eintragen."),
        (None, None),
        ("Hinweise zur Vorlage", None),
        ("Herkunft", "Prüfpunkte 1–39 der Vorlage „Auditcheckliste_allg“ (Spalte „Quelle“ = Vorlage-Nr.). "
                     "Die englischen Texte stehen in der ausgeblendeten Spalte Q."),
        ("Verfügbarkeit", "Die Vorlage nannte 90 % (DE) bzw. 80 % (EN) technische Verfügbarkeit. Zielwert auf „Einstellungen“ festlegen."),
        (None, None),
        ("Quellen", None),
        ("DIN 31051", "Grundlagen der Instandhaltung – https://www.baunormenlexikon.de/norm/din-31051/8d1293a4-de10-49d4-9fee-0b415e6545a1"),
        ("DIN EN 13306", "Begriffe der Instandhaltung – https://ihb.tuev-media.de/docs/din_en_13306.html"),
        ("DIN EN 15341", "Leistungskennzahlen der Instandhaltung – https://blog.ccc-industriesoftware.de/kennzahlen-und-kpis-instandhaltung/"),
        ("IATF 16949, 8.5.1.5", "TPM, Schlüsselanlagen, Ersatzteile, Konservierung – https://community.advisera.com/topic/total-productive-maintenance-in-iatf-16949/"),
        ("BetrSichV", "§ 10 Instandhaltung, § 14 Prüfung – https://www.gesetze-im-internet.de/betrsichv_2015/BJNR004910015.html"),
        ("TRBS 1112", "Instandhaltung – https://www.baua.de/DE/Angebote/Regelwerk/TRBS/pdf/TRBS-1112.pdf"),
        ("DGUV Vorschrift 3", "Prüfung elektrischer Anlagen – https://www.forum-verlag.com/fachwissen/elektrosicherheit-und-elektrotechnik/prueffristen-nach-dguv-v3/"),
        ("AwSV / WGK", "https://www.umweltbundesamt.de/system/files/medien/421/dokumente/info_awsv_dieter_wgk.pdf"),
    ]
    r = 3
    for a, b in rows:
        if a and b is None:
            ws.cell(r, 1, a).font = f(11, True, NAVY)
        elif a:
            ws.cell(r, 1, a).font = f(10, True)
            ws.cell(r, 2, b).font = f()
            ws.cell(r, 2).alignment = WRAP
            ws.cell(r, 1).alignment = Alignment(vertical="top")
        r += 1
    if beispiel:
        ws.cell(r + 1, 1, "Beispieldatei").font = f(11, True, ROT_TXT)
        ws.cell(r + 1, 2, "Diese Datei enthält erfundene Beispielbewertungen, damit das Dashboard etwas zeigt. "
                          "Für ein echtes Audit die Datei „…_Vorlage.xlsx“ verwenden.").font = f(10, color=ROT_TXT)

    ws_c["G4"].comment = Comment("Gewichteter Erfüllungsgrad – Details im Dashboard.", "Audit")
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


if __name__ == "__main__":
    build(False, os.path.join(HERE, "Audit_IH_Allgemein_Vorlage.xlsx"))
    build(True, os.path.join(HERE, "Audit_IH_Allgemein_Beispiel.xlsx"))
    print("fertig")
